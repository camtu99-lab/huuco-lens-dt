import base64
import csv
import html
import io
import json
import math
import random
import re
import time
import unicodedata
from datetime import datetime

import streamlit as st
import streamlit.components.v1 as components
from google import genai
from google.genai import errors, types
from PIL import Image
from collections import Counter

from rdkit import Chem, RDLogger
from rdkit.Chem import Descriptors, Draw, rdMolDescriptors
from rdkit.Chem.EnumerateStereoisomers import EnumerateStereoisomers, StereoEnumerationOptions

RDLogger.DisableLog("rdApp.*")

MODEL = st.secrets.get("GEMINI_MODEL", "gemini-2.5-flash")
MAX_PER_SESSION = int(st.secrets.get("MAX_PER_SESSION", 15))
BENCH_MAX = int(st.secrets.get("BENCH_MAX", 30))

APP_NAME = st.secrets.get("APP_NAME", "MolSnap")
APP_ICON = st.secrets.get("APP_ICON", "🧪")

st.set_page_config(page_title=APP_NAME, page_icon=APP_ICON, layout="wide", initial_sidebar_state="auto")

RULES = """Bạn là giáo viên Hóa hữu cơ Việt Nam kiêm chuyên gia nhận dạng công thức. Quy ước trong SGK Việt Nam: công thức khai triển/thu gọn (CH3-CH2-OH, CH3COOH), công thức khung (mỗi đầu mút/đỉnh là C, H gắn vào C không vẽ), vòng benzen vẽ lục giác với 3 liên kết đôi xen kẽ hoặc có vòng tròn bên trong, C6H5- hoặc Ph là phenyl, nhóm chức như -OH, -CHO, -COOH, -COO-, -CO-, -NH2, -NO2, -C≡N, -X. Đếm cẩn thận số C, vị trí nhánh và nhóm thế trên vòng (o-, m-, p-). Danh pháp dùng cách đọc tiếng Việt theo SGK (metan, etanol, propan-2-ol, axit etanoic, etanal, benzen, phenol, anilin, toluen, etyl axetat...), kèm tên thường gọi nếu có."""

FMT = """Chỉ trả về MỘT đối tượng JSON, không markdown, không lời dẫn, các khóa:
{"smiles":"SMILES chuẩn (không có H tường minh)","do_tin_cay":số 0-1,"ten_danh_phap":"tên thay thế/IUPAC đọc kiểu SGK Việt Nam","ten_thuong_goi":"tên thông thường hoặc rỗng","loai_hop_chat":"ví dụ: ancol no đơn chức","nhom_chuc":["..."],"trang_thai_nhiet_do":"thể ở điều kiện thường, nhận xét nhiệt độ sôi/tan","tinh_chat_vat_ly":["..."],"tinh_chat_hoa_hoc":["phản ứng đặc trưng kèm phương trình ngắn, điều kiện"],"dieu_che_ung_dung":["..."],"dong_phan":"nhận xét ngắn nếu có ý nghĩa","luu_y":["điều chưa chắc chắn trong ảnh, nét vẽ mơ hồ, cảnh báo an toàn"]}
Nếu ảnh không phải công thức hóa học hoặc không đọc được, đặt "smiles":"" và giải thích trong luu_y."""

SUB = str.maketrans("0123456789", "₀₁₂₃₄₅₆₇₈₉")


@st.cache_resource
def get_client():
    return genai.Client(api_key=st.secrets["GEMINI_API_KEY"])


def to_jpeg(file_bytes: bytes) -> bytes:
    img = Image.open(io.BytesIO(file_bytes)).convert("RGB")
    img.thumbnail((1600, 1600))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=90)
    return buf.getvalue()


def ask_ai(contents, temperature=None):
    """Gọi Gemini, tự thử lại khi máy chủ quá tải (503) hoặc hết lượt tạm thời (429),
    rồi chuyển sang mô hình dự phòng."""
    models = [MODEL]
    fb = st.secrets.get("GEMINI_FALLBACK_MODEL", "gemini-3.5-flash-lite")
    if fb and fb != MODEL:
        models.append(fb)
    last = None
    for m in models:
        for i in range(3):
            try:
                resp = get_client().models.generate_content(
                    model=m,
                    contents=contents,
                    config=types.GenerateContentConfig(
                        response_mime_type="application/json",
                        **({} if temperature is None else {"temperature": temperature}),
                    ),
                )
                text = (resp.text or "").replace("```json", "").replace("```", "").strip()
                return json.loads(text[text.index("{"): text.rindex("}") + 1])
            except errors.ServerError as e:
                last = e
                time.sleep(2 * (i + 1))
            except errors.ClientError as e:
                code = getattr(e, "code", None)
                if code == 429:
                    last = e
                    time.sleep(3 * (i + 1))
                elif code == 404:
                    last = e
                    break  # mô hình này không còn dùng được, chuyển sang mô hình tiếp theo
                else:
                    raise
    raise last


def analyze_image(file_bytes: bytes):
    prompt = (
        RULES
        + "\n\nĐây là ảnh chụp công thức cấu tạo viết tay. Hãy nhận dạng đúng một chất hữu cơ trong ảnh. "
        + "Chép ĐÚNG như hình vẽ, KHÔNG tự sửa lỗi hóa trị và không thêm hoặc bớt nguyên tử. "
        + "Nếu hình vẽ sai (ví dụ nguyên tử C có 5 liên kết) vẫn ghi SMILES đúng theo hình và nêu rõ lỗi trong luu_y.\n"
        + FMT
    )
    return ask_ai([types.Part.from_bytes(data=to_jpeg(file_bytes), mime_type="image/jpeg"), prompt], temperature=0)


def analyze_smiles(smiles: str):
    prompt = (
        RULES
        + f"\n\nChất có SMILES: {smiles}\nGiữ nguyên SMILES này trong khóa \"smiles\", đặt do_tin_cay=1 và cung cấp tên cùng tính chất.\n"
        + FMT
    )
    return ask_ai([prompt])


def chem_numbers(smiles: str):
    """Tính trực tiếp bằng RDKit, không phụ thuộc AI."""
    mol = Chem.MolFromSmiles(smiles)
    if mol is None:
        return None, None
    formula = rdMolDescriptors.CalcMolFormula(mol)
    pretty = "".join(c.translate(SUB) if c.isdigit() else c for c in formula)
    molh = Chem.AddHs(mol)
    cnt = {}
    for a in molh.GetAtoms():
        cnt[a.GetSymbol()] = cnt.get(a.GetSymbol(), 0) + 1
    dbe = None
    if cnt.get("C"):
        x = sum(cnt.get(k, 0) for k in ("F", "Cl", "Br", "I"))
        dbe = (2 * cnt["C"] + 2 + cnt.get("N", 0) + cnt.get("P", 0) - cnt.get("H", 0) - x) / 2
    return mol, {"formula": pretty, "mass": Descriptors.MolWt(mol), "dbe": dbe}


def bullets(title, items):
    if items:
        st.markdown(f"**{title}**")
        for t in [items] if isinstance(items, str) else items:
            st.markdown(f"- {t}")


def render(r: dict, from_image: bool):
    smiles = (r.get("smiles") or "").strip()
    if not smiles:
        st.warning("Chưa nhận dạng được công thức trong ảnh.")
    mol, num = chem_numbers(smiles) if smiles else (None, None)
    if smiles and mol is None:
        st.error("Mã SMILES không hợp lệ. Hãy sửa ở ô bên dưới rồi bấm Phân tích lại.")
    c1, c2 = st.columns(2)
    with c1:
        if mol is not None:
            st.image(Draw.MolToImage(mol, size=(420, 300)))
    with c2:
        if smiles:
            st.markdown("**SMILES**")
            st.code(smiles, language=None)
        if r.get("ten_danh_phap"):
            st.markdown(f"**Tên (danh pháp):** {r['ten_danh_phap']}")
        if r.get("ten_thuong_goi"):
            st.markdown(f"**Tên thường gọi:** {r['ten_thuong_goi']}")
        if r.get("loai_hop_chat"):
            st.markdown(f"**Loại hợp chất:** {r['loai_hop_chat']}")
        if num:
            st.markdown(f"**Công thức phân tử:** {num['formula']}")
            st.markdown(f"**Khối lượng mol:** {num['mass']:.2f} g/mol")
            if num["dbe"] is not None:
                k = num["dbe"]
                st.markdown(f"**Độ bất bão hòa (k):** {int(k) if k == int(k) else k}")
    if from_image and isinstance(r.get("do_tin_cay"), (int, float)):
        p = max(0.0, min(1.0, float(r["do_tin_cay"])))
        st.progress(p, text=f"Độ tin cậy nhận dạng: {round(p * 100)}%")
    if r.get("nhom_chuc"):
        st.markdown("**Nhóm chức:** " + " · ".join(r["nhom_chuc"]))
    bullets("Trạng thái, nhiệt độ", r.get("trang_thai_nhiet_do"))
    bullets("Tính chất vật lí", r.get("tinh_chat_vat_ly"))
    bullets("Tính chất hóa học", r.get("tinh_chat_hoa_hoc"))
    bullets("Điều chế và ứng dụng", r.get("dieu_che_ung_dung"))
    bullets("Đồng phân", r.get("dong_phan"))
    if r.get("luu_y"):
        st.info("Lưu ý: " + " • ".join([r["luu_y"]] if isinstance(r["luu_y"], str) else r["luu_y"]))



# ---------------- Phản ứng hóa học ----------------
RULES_RX = """Bạn là giáo viên Hóa học Việt Nam. Hãy đọc phương trình phản ứng (viết tay hoặc gõ) và phân loại theo cách của SGK Việt Nam (Hóa 8-12). Các loại thường gặp: phản ứng thế (thế halogen, thế nitro, thế brom vào phenol...), cộng (cộng H2, Br2, HX, H2O), tách (tách H2O, tách HX, tách H2), trùng hợp, trùng ngưng, oxi hóa hoàn toàn (đốt cháy), oxi hóa không hoàn toàn, oxi hóa - khử, hóa hợp, phân hủy, trao đổi, trung hòa, este hóa, thủy phân, xà phòng hóa, tráng bạc, lên men, cracking, đề hiđro hóa, hiđro hóa. Một phản ứng có thể thuộc nhiều loại cùng lúc (ví dụ: cộng đồng thời là oxi hóa - khử); hãy nêu loại chính trước. Đọc kỹ hệ số và điều kiện (nhiệt độ, xúc tác, áp suất). Với muối/bazơ vô cơ dùng SMILES dạng ion, ví dụ NaOH là [Na+].[OH-], HCl là Cl, H2O là O, O2 là O=O, H2 là [H][H]."""

RFMT = """Chỉ trả về MỘT đối tượng JSON, không markdown, không lời dẫn:
{"phuong_trinh":"phương trình đọc được, viết dạng chữ, dùng -> giữa hai vế",
"chat_tham_gia":[{"he_so":số nguyên,"ten":"tên chất","cong_thuc":"công thức phân tử","smiles":"SMILES"}],
"san_pham":[{"he_so":số nguyên,"ten":"","cong_thuc":"","smiles":""}],
"dieu_kien":"nhiệt độ, xúc tác... hoặc rỗng",
"loai_phan_ung":["loại chính","loại phụ nếu có"],
"giai_thich":"2-4 câu: dựa vào sự thay đổi nào của cấu tạo/nhóm chức/số oxi hóa mà kết luận loại phản ứng",
"phuong_trinh_can_bang":"phương trình đã cân bằng đúng nếu ảnh chưa cân bằng hoặc sai, ngược lại để rỗng",
"luu_y":["điều chưa chắc chắn khi đọc ảnh, ví dụ hệ số hoặc chất bị mờ"]}
Nếu không phải phương trình phản ứng hoặc không đọc được, trả về chat_tham_gia và san_pham rỗng và giải thích trong luu_y."""


def analyze_reaction(img_bytes, text):
    head = RULES_RX + "\n\n"
    if img_bytes:
        prompt = head + "Đây là ảnh chụp một phương trình phản ứng hóa học.\n" + RFMT
        return ask_ai([types.Part.from_bytes(data=to_jpeg(img_bytes), mime_type="image/jpeg"), prompt])
    return ask_ai([head + f"Phương trình: {text}\n" + RFMT])


def atom_balance(species):
    """Đếm nguyên tử và điện tích mỗi vế bằng RDKit. Trả về None nếu có SMILES lỗi."""
    tot, charge = {}, 0
    for coef, smi in species:
        mol = Chem.MolFromSmiles(smi or "")
        if mol is None:
            return None
        for a in Chem.AddHs(mol).GetAtoms():
            tot[a.GetSymbol()] = tot.get(a.GetSymbol(), 0) + coef
            charge += coef * a.GetFormalCharge()
    return tot, charge


def to_int(v):
    try:
        return max(1, int(v))
    except Exception:
        return 1


def render_reaction(r: dict):
    left, right = r.get("chat_tham_gia") or [], r.get("san_pham") or []
    if r.get("phuong_trinh"):
        st.markdown("**Phương trình đọc được**")
        st.code(r["phuong_trinh"], language=None)
    types_ = r.get("loai_phan_ung") or []
    if types_:
        st.success("**Loại phản ứng:** " + " · ".join(types_))
    if r.get("giai_thich"):
        st.markdown("**Vì sao:** " + r["giai_thich"])
    if r.get("dieu_kien"):
        st.markdown("**Điều kiện:** " + r["dieu_kien"])

    rows, mols, legends = [], [], []
    for role, lst in (("Tham gia", left), ("Sản phẩm", right)):
        for sp in lst:
            m = Chem.MolFromSmiles(sp.get("smiles") or "")
            mass = f"{Descriptors.MolWt(m):.2f}" if m is not None else "—"
            rows.append({"Vai trò": role, "Hệ số": to_int(sp.get("he_so")), "Tên": sp.get("ten", ""),
                         "Công thức": sp.get("cong_thuc", ""), "SMILES": sp.get("smiles", ""), "M (g/mol)": mass})
            if m is not None and m.GetNumHeavyAtoms() > 1:
                mols.append(m)
                legends.append(sp.get("ten") or sp.get("cong_thuc") or "")
    if rows:
        st.dataframe(rows, hide_index=True, width="stretch")
    if mols:
        try:
            st.image(Draw.MolsToGridImage(mols, molsPerRow=3, subImgSize=(260, 190), legends=legends))
        except Exception:
            pass

    if left and right:
        bl = atom_balance([(to_int(s.get("he_so")), s.get("smiles")) for s in left])
        br = atom_balance([(to_int(s.get("he_so")), s.get("smiles")) for s in right])
        if bl is None or br is None:
            st.info("Không kiểm tra được cân bằng vì có chất có SMILES không hợp lệ.")
        elif bl == br:
            st.success("Phương trình **đã cân bằng** (RDKit đếm nguyên tử và điện tích hai vế bằng nhau).")
        else:
            diff = [f"{k}: {bl[0].get(k, 0)} ≠ {br[0].get(k, 0)}" for k in sorted(set(bl[0]) | set(br[0]))
                    if bl[0].get(k, 0) != br[0].get(k, 0)]
            if bl[1] != br[1]:
                diff.append(f"điện tích: {bl[1]} ≠ {br[1]}")
            st.warning("Phương trình **chưa cân bằng** theo hệ số đọc được. Lệch: " + "; ".join(diff))
            if r.get("phuong_trinh_can_bang"):
                st.markdown("**Gợi ý cân bằng:**")
                st.code(r["phuong_trinh_can_bang"], language=None)
    if r.get("luu_y"):
        st.info("Lưu ý: " + " • ".join([r["luu_y"]] if isinstance(r["luu_y"], str) else r["luu_y"]))


def reaction_ui():
    top = st.container()  # lỗi và lý do sai hiện ở đầu tab
    src = st.radio("Nguồn phương trình", ["🖼️ Tải ảnh", "📷 Chụp ảnh", "⌨️ Gõ phương trình"], horizontal=True, key="rx_src")
    img, txt = None, None
    if src.startswith("🖼️"):
        f = st.file_uploader("Chọn ảnh phương trình", type=["jpg", "jpeg", "png", "webp"], key="rx_up")
        img = f.getvalue() if f else None
    elif src.startswith("📷"):
        f = st.camera_input("Chụp phương trình", key="rx_cam")
        img = f.getvalue() if f else None
    else:
        txt = st.text_input("Phương trình", placeholder="Ví dụ: CH2=CH2 + Br2 -> CH2Br-CH2Br", key="rx_txt").strip()
    if st.button("Nhận dạng và phân loại phản ứng", type="primary", disabled=not (img or txt), key="rx_go"):
        if st.session_state.n >= MAX_PER_SESSION:
            show_top_error(top, "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục.")
        else:
            with st.spinner("Đang phân tích…"):
                try:
                    st.session_state.rx = analyze_reaction(img, txt)
                    add_history(((st.session_state.rx or {}).get('phuong_trinh') or txt or 'Phương trình')[:60], 'Phản ứng')
                    st.session_state.n += 1
                except Exception as e:
                    st.session_state.rx = None
                    show_top_error(top, "Có lỗi khi phân tích, hãy thử lại hoặc chụp rõ hơn.", e)
    if st.session_state.get("rx"):
        st.divider()
        render_reaction(st.session_state.rx)



# ---------------- Đồng phân ----------------
RULES_ISO = """Bạn là giáo viên Hóa hữu cơ Việt Nam. Xác định đồng phân theo SGK: đồng phân cấu tạo (mạch cacbon, vị trí nhóm chức hoặc liên kết bội, nhóm chức) và đồng phân lập thể (hình học cis-trans, quang học). Tên gọi dùng danh pháp tiếng Việt theo SGK (butan-1-ol, đimetyl ete, 2-metylpropan-2-ol...)."""

PRETTY = str.maketrans("0123456789", "₀₁₂₃₄₅₆₇₈₉")


def formula_of(m):
    return rdMolDescriptors.CalcMolFormula(m)


def pretty(f):
    return "".join(c.translate(PRETTY) if c.isdigit() else c for c in f)


def noiso(m):
    return Chem.MolToSmiles(m, isomericSmiles=False)


def stereo_count(m):
    try:
        opts = StereoEnumerationOptions(onlyUnassigned=True, unique=True, tryEmbedding=False, maxIsomers=32)
        return max(1, len(list(EnumerateStereoisomers(m, options=opts))))
    except Exception:
        return 1


def mirror(smi):
    return smi.replace("@@", "\x00").replace("@", "@@").replace("\x00", "@")


def relation(m1, m2):
    f1, f2 = formula_of(m1), formula_of(m2)
    if f1 != f2:
        return "khac_cong_thuc"
    if noiso(m1) != noiso(m2):
        return "cau_tao"
    i1, i2 = Chem.MolToSmiles(m1), Chem.MolToSmiles(m2)
    if i1 == i2:
        return "cung_chat"
    mir = Chem.MolFromSmiles(mirror(i1))
    if mir is not None and Chem.MolToSmiles(mir) == i2:
        return "doi_quang"
    if not Chem.FindMolChiralCenters(m1) and not Chem.FindMolChiralCenters(m2):
        return "hinh_hoc"
    return "lap_the"


# ---- Sinh đồng phân cấu tạo bằng thuật toán (không dùng AI) ----
VALENCE = {"C": 4, "N": 3, "O": 2, "S": 2, "F": 1, "Cl": 1, "Br": 1, "I": 1}
_BT = {1: Chem.BondType.SINGLE, 2: Chem.BondType.DOUBLE, 3: Chem.BondType.TRIPLE}
ISO_TIME = float(st.secrets.get("ISO_TIME", 15))


class TooBig(Exception):
    pass


def _canon(atoms, edges, strict=False):
    rw = Chem.RWMol()
    for e in atoms:
        rw.AddAtom(Chem.Atom(e))
    for i, j, o in edges:
        rw.AddBond(i, j, _BT[o])
    m = rw.GetMol()
    Chem.SanitizeMol(m)
    if strict:
        for b in m.GetBonds():
            if b.GetBondType() == Chem.BondType.TRIPLE and b.IsInRing():
                raise ValueError("liên kết ba trong vòng")
    return Chem.MolToSmiles(m, isomericSmiles=False)


def enumerate_isomers(counts, nH, time_limit=15, max_graphs=200000, strict=False):
    """Liệt kê TẤT CẢ đồng phân cấu tạo của công thức (mọi đồ thị liên thông đúng hóa trị, không trùng nhau).
    strict=True: chỉ mạch hở hoặc 1 vòng, không có O–O, không liên kết ba trong vòng."""
    els = sorted(counts)
    N = sum(counts.values())
    sigma = sum(VALENCE[e] * counts[e] for e in els)
    if sigma - nH < 0 or (sigma - nH) % 2:
        return []
    B = (sigma - nH) // 2
    if B < N - 1:
        return []
    t0, memo, total = time.time(), {}, [0]

    def build(sub):
        if sub in memo:
            return memo[sub]
        n = sum(sub)
        res = {}
        if n == 1:
            e = els[[i for i, c in enumerate(sub) if c][0]]
            res[_canon([e], [])] = ([e], [])
            memo[sub] = res
            return res
        upper = B - (N - n)
        for i, c in enumerate(sub):
            if not c:
                continue
            prev = list(sub)
            prev[i] -= 1
            e, vn = els[i], VALENCE[els[i]]
            for atoms, edges in list(build(tuple(prev)).values()):
                T = sum(o for _, _, o in edges)
                if T + 1 > upper:
                    continue
                used = [0] * len(atoms)
                for a, b, o in edges:
                    used[a] += o
                    used[b] += o
                free = [VALENCE[atoms[k]] - used[k] for k in range(len(atoms))]
                k = len(atoms)

                def rec(j, rem_new, added, att):
                    if time.time() - t0 > time_limit:
                        raise TooBig("time")
                    if j == k:
                        if att and T + added <= upper and (n < N or T + added == B):
                            na, ne = atoms + [e], edges + [(a, k, o) for a, o in att]
                            if strict:
                                if len(ne) - len(na) + 1 > 1:
                                    return
                                if e == "O" and any(atoms[a] == "O" for a, _ in att):
                                    return
                            try:
                                key = _canon(na, ne, strict)
                            except Exception:
                                return
                            if key not in res:
                                res[key] = (na, ne)
                                total[0] += 1
                                if total[0] > max_graphs:
                                    raise TooBig("size")
                        return
                    rec(j + 1, rem_new, added, att)
                    for o in (1, 2, 3):
                        if o <= rem_new and o <= free[j] and T + added + o <= upper:
                            rec(j + 1, rem_new - o, added + o, att + [(j, o)])

                rec(0, vn, 0, [])
        memo[sub] = res
        return res

    return sorted(build(tuple(counts[e] for e in els)).keys())


def parse_formula(s):
    s = (s or "").translate(str.maketrans("₀₁₂₃₄₅₆₇₈₉", "0123456789")).replace(" ", "")
    toks = re.findall(r"([A-Z][a-z]?)(\d*)", s)
    if not toks or "".join(a + b for a, b in toks) != s:
        raise ValueError("không phải công thức phân tử")
    c = {}
    for a, b in toks:
        c[a] = c.get(a, 0) + (int(b) if b else 1)
    nH = c.pop("H", 0)
    bad = [e for e in c if e not in VALENCE]
    if bad or not c:
        raise ValueError("nguyên tố chưa hỗ trợ: " + ", ".join(bad or ["?"]))
    return c, nH


def is_simple_formula(s):
    s = (s or "").translate(str.maketrans("₀₁₂₃₄₅₆₇₈₉", "0123456789")).replace(" ", "")
    toks = re.findall(r"([A-Z][a-z]?)(\d*)", s)
    if not toks or "".join(a + b for a, b in toks) != s:
        return False
    syms = [a for a, _ in toks]
    return len(syms) == len(set(syms)) and (any(b for _, b in toks) or "H" in syms)


def hill_key(counts, nH):
    parts = []
    if "C" in counts:
        parts.append(("C", counts["C"]))
    if nH:
        parts.append(("H", nH))
    parts += [(e, counts[e]) for e in sorted(counts) if e != "C"]
    return "".join(f"{e}{n if n > 1 else ''}" for e, n in parts)


def mol_counts(m):
    c = atom_counts(m)
    nH = c.pop("H", 0)
    return c, nH


# Phân loại theo chương trình phổ thông (dùng riêng cho tab đồng phân)
ISO_GROUPS = [
    ("ancol bậc I", "[CX4;H2,H3][OX2H]"), ("ancol bậc II", "[CX4;H1]([#6])([#6])[OX2H]"),
    ("ancol bậc III", "[CX4;H0]([#6])([#6])([#6])[OX2H]"), ("phenol", "c[OX2H]"),
    ("ete", "[#6;!$([#6]=O)][OX2][#6;!$([#6]=O)]"), ("anđehit", "[CX3;H1,H2;!$(C[OX2,NX3])]=O"),
    ("xeton", "[#6][CX3](=O)[#6]"), ("axit cacboxylic", "[CX3](=O)[OX2H1]"), ("este", "[CX3](=O)[OX2][#6]"),
    ("amin bậc I", "[NX3;H2;!$(NC=[O,N,S])][#6]"), ("amin bậc II", "[NX3;H1;!$(NC=[O,N,S])]([#6])[#6]"),
    ("amin bậc III", "[NX3;H0;!$(NC=[O,N,S])]([#6])([#6])[#6]"), ("amit", "[NX3][CX3]=O"),
    ("nitro", "[N+](=O)[O-]"), ("nitrin", "[CX2]#[NX1]"), ("anken (C=C)", "[CX3]=[CX3]"), ("ankin (C≡C)", "[CX2]#[CX2]"),
    ("hợp chất thơm", "a"), ("dẫn xuất halogen", "[#6][F,Cl,Br,I]"),
]
_ISO_GP = [(n, Chem.MolFromSmarts(p)) for n, p in ISO_GROUPS]
UNSTABLE = [
    ("enol (OH trên C=C)", "[OX2H][CX3]=[CX3]"), ("ynol (OH trên C≡C)", "[OX2H][CX2]#[CX2]"),
    ("hai nhóm OH trên cùng một C", "[CX4]([OX2H])[OX2H]"), ("hemiaxetal", "[CX4]([OX2H])[OX2][#6]"),
    ("OH và halogen trên cùng một C", "[CX4]([OX2H])[F,Cl,Br,I]"), ("OH và N trên cùng một C", "[CX4]([OX2H])[NX3]"),
]
_UNST = [(n, Chem.MolFromSmarts(p)) for n, p in UNSTABLE]


_PAIR = Chem.MolFromSmarts("[CX3](=O)[OX2,NX3]")
_NITRO = Chem.MolFromSmarts("[N+](=O)[O-]")


def func_units(m):
    """Số nhóm chức chứa dị tố: axit/este/amit và nitro mỗi nhóm chỉ tính 1 dù có nhiều nguyên tử."""
    n_het = sum(1 for a in m.GetAtoms() if a.GetSymbol() in ("O", "N", "S", "F", "Cl", "Br", "I"))
    return n_het - len(m.GetSubstructMatches(_PAIR)) - 2 * len(m.GetSubstructMatches(_NITRO))


def is_mono_open(m):
    """Đơn chức (tối đa 1 nhóm chức chứa dị tố) và không có vòng no/không thơm (vòng benzen vẫn được giữ)."""
    return func_units(m) <= 1 and not any(a.IsInRing() and not a.GetIsAromatic() for a in m.GetAtoms())


def iso_labels(m):
    lab = [n for n, p in _ISO_GP if m.HasSubstructMatch(p)]
    has_ring = any(not a.GetIsAromatic() and a.IsInRing() for a in m.GetAtoms())
    if not lab:
        if all(a.GetSymbol() == "C" for a in m.GetAtoms()):
            return ["xicloankan"] if has_ring else ["ankan"]
        lab = ["khác"]
    if has_ring:
        lab.append("có vòng không thơm")
    return lab


@st.cache_data(show_spinner=False, max_entries=64)
def iso_enumerate(fkey, strict):
    counts, nH = parse_formula(fkey)
    try:
        return {"ok": True, "smiles": enumerate_isomers(counts, nH, time_limit=ISO_TIME, strict=strict)}
    except TooBig as e:
        return {"ok": False, "why": str(e)}


def stereo_split(m):
    """(tổng, hình học, quang học): số dạng lập thể; hình học = cis/trans bỏ qua quang học; quang học = bỏ qua cis/trans."""
    try:
        opts = StereoEnumerationOptions(onlyUnassigned=True, unique=True, tryEmbedding=False, maxIsomers=64)
        isos = list(EnumerateStereoisomers(m, options=opts))
    except Exception:
        return 1, 1, 1
    geo, opt = set(), set()
    for x in isos:
        a = Chem.Mol(x)  # hình học: giữ cis/trans của nối đôi và cấu hình tương đối trên vòng; bỏ tâm bất đối ngoài vòng và gộp cặp đối quang
        for at in a.GetAtoms():
            if not at.IsInRing():
                at.SetChiralTag(Chem.ChiralType.CHI_UNSPECIFIED)
        s1 = Chem.MolToSmiles(a)
        mir = Chem.MolFromSmiles(mirror(s1))
        geo.add(min(s1, Chem.MolToSmiles(mir)) if mir is not None else s1)
        b = Chem.Mol(x)  # quang học: giữ mọi tâm bất đối, bỏ cis/trans của nối đôi
        for bd in b.GetBonds():
            if bd.GetBondType() == Chem.BondType.DOUBLE:
                bd.SetStereo(Chem.BondStereo.STEREONONE)
                bd.SetBondDir(Chem.BondDir.NONE)
        opt.add(Chem.MolToSmiles(b))
    return max(1, len(isos)), max(1, len(geo)), max(1, len(opt))


# Từ vựng nhóm mà AI được phép dùng; mỗi từ mở rộng thành các nhãn do bộ máy phân loại
GROUP_VOCAB = {
    "ankan": ["ankan"], "xicloankan": ["xicloankan"], "anken": ["anken (C=C)"], "ankin": ["ankin (C≡C)"],
    "thơm": ["hợp chất thơm"], "ancol": ["ancol bậc I", "ancol bậc II", "ancol bậc III"],
    "ancol bậc I": ["ancol bậc I"], "ancol bậc II": ["ancol bậc II"], "ancol bậc III": ["ancol bậc III"],
    "phenol": ["phenol"], "ete": ["ete"], "anđehit": ["anđehit"], "xeton": ["xeton"],
    "axit cacboxylic": ["axit cacboxylic"], "este": ["este"],
    "amin": ["amin bậc I", "amin bậc II", "amin bậc III"], "amin bậc I": ["amin bậc I"],
    "amin bậc II": ["amin bậc II"], "amin bậc III": ["amin bậc III"], "amit": ["amit"], "nitro": ["nitro"],
    "nitrin": ["nitrin"], "dẫn xuất halogen": ["dẫn xuất halogen"],
}
PRESETS = {
    "mono": {"ben": True, "chuc": "don", "mach": "ho_hoac_thom", "nhom": [], "nhom_ten": [], "kieu": "bat_ky", "lap_the": "khong", "defaults": []},
    "pho": {"ben": True, "chuc": "bat_ky", "mach": "bat_ky", "nhom": [], "nhom_ten": [], "kieu": "bat_ky", "lap_the": "khong", "defaults": []},
    "math": {"ben": False, "chuc": "bat_ky", "mach": "bat_ky", "nhom": [], "nhom_ten": [], "kieu": "bat_ky", "lap_the": "khong", "defaults": []},
}


def norm_params(raw):
    """Kiểm tra tham số AI trích ra: chỉ nhận giá trị trong danh mục. Giá trị AI để trống thì dùng mặc định và ghi rõ."""
    raw = raw or {}
    warn, defaults = [], []
    names = []
    for g in raw.get("nhom") or []:
        key = str(g).strip().lower().replace("đ", "đ")
        hit = next((k for k in GROUP_VOCAB if k.lower() == key), None)
        if hit:
            names.append(hit)
        else:
            warn.append(f"Bỏ qua nhóm không có trong danh mục: {g}")
    p = {"ben": True, "nhom_ten": names, "nhom": [GROUP_VOCAB[n] for n in names]}
    kieu = raw.get("kieu_nhom")
    p["kieu"] = kieu if kieu in ("bat_ky", "tat_ca") else "bat_ky"
    for k, ok, dflt in (("chuc", ("don", "da", "bat_ky"), "don"), ("mach", ("ho", "vong", "bat_ky"), "ho_hoac_thom"),
                        ("lap_the", ("khong", "hinh_hoc", "quang_hoc", "tat_ca"), "khong")):
        v = raw.get(k)
        if v in ok:
            p[k] = v
        else:
            p[k] = dflt
            defaults.append(k)
    if "xicloankan" in names and "mach" in defaults:
        p["mach"] = "bat_ky"
    p["defaults"] = defaults
    return p, warn


def pass_filters(m, labels, p):
    u = func_units(m)
    if (p["chuc"] == "don" and u > 1) or (p["chuc"] == "da" and u < 2):
        return False
    rings = m.GetRingInfo().NumRings()
    nonarom = any(a.IsInRing() and not a.GetIsAromatic() for a in m.GetAtoms())
    if (p["mach"] == "ho" and rings) or (p["mach"] == "vong" and not rings) or (p["mach"] == "ho_hoac_thom" and nonarom):
        return False
    if p["nhom"]:
        hit = [bool(set(labels) & set(g)) for g in p["nhom"]]
        return all(hit) if p["kieu"] == "tat_ca" else any(hit)
    return True


PARSE_PROMPT = """Bạn là bộ PHÂN TÍCH CÂU HỎI Hóa học. Nhiệm vụ DUY NHẤT: trích tham số từ câu hỏi về đồng phân. TUYỆT ĐỐI không đếm, không liệt kê chất, không trả lời câu hỏi.
Từ vựng nhóm (chỉ dùng đúng các chuỗi này): ankan, xicloankan, anken, ankin, thơm, ancol, ancol bậc I, ancol bậc II, ancol bậc III, phenol, ete, anđehit, xeton, axit cacboxylic, este, amin, amin bậc I, amin bậc II, amin bậc III, amit, nitro, nitrin, dẫn xuất halogen.
Chỉ trả về MỘT đối tượng JSON, không markdown:
{"cong_thuc":"công thức phân tử viết liền, ví dụ C4H8O2, hoặc null",
"chat_goc":"tên hoặc SMILES của chất nếu câu hỏi nhắc một chất cụ thể thay vì công thức, hoặc null",
"nhom":["nhóm mà đồng phân cần tìm phải thuộc, lấy từ từ vựng; để [] nếu câu hỏi không nêu nhóm"],
"kieu_nhom":"bat_ky nếu chỉ cần thuộc ít nhất một nhóm, tat_ca nếu phải có đồng thời mọi nhóm, null nếu không nêu",
"chuc":"don nếu câu hỏi nói đơn chức, da nếu nói đa chức hoặc tạp chức, bat_ky nếu nói rõ tính mọi loại, null nếu không nhắc",
"mach":"ho nếu mạch hở, vong nếu mạch vòng, bat_ky nếu nói rõ tính cả hai, null nếu không nhắc",
"lap_the":"khong nếu nói không tính đồng phân lập thể, hinh_hoc nếu tính thêm đồng phân hình học (cis-trans), quang_hoc nếu tính thêm đồng phân quang học, tat_ca nếu tính mọi đồng phân lập thể, null nếu không nhắc",
"can_hoi_lai":"một câu hỏi làm rõ nếu thiếu công thức hoặc chất, hoặc câu hỏi mơ hồ; nếu không thì để rỗng",
"tom_tat":"câu hỏi được hiểu lại ngắn gọn bằng tiếng Việt"}
Không đoán những gì câu hỏi không nói: để null hoặc []. Ví dụ: 'Có bao nhiêu đồng phân este của C4H8O2?' thì nhom=["este"], cong_thuc="C4H8O2", các mục còn lại null."""


def parse_question(q):
    raw = ask_ai([PARSE_PROMPT + "\n\nCâu hỏi: " + q], temperature=0)
    p, warn = norm_params(raw)
    return p, warn, raw


def resolve_formula(img, txt):
    """Trả về (counts, nH, SMILES chất gốc hoặc None, lỗi hoặc None)."""
    if img:
        smi, _ = transcribe_student(img)
        m, err = parse_student(smi)
        if m is None:
            return None, None, None, "Hình vẽ sai nên không xác định được công thức: " + (err or "không đọc được hình")
    elif is_simple_formula(txt):
        try:
            c, h = parse_formula(txt)
            return c, h, None, None
        except ValueError as e:
            return None, None, None, f"Công thức không hợp lệ ({e})."
    else:
        m, err = parse_student(txt) if smiles_like(txt) else (None, None)
        if m is None:
            smi = to_smiles(None, txt)  # tên chất hoặc công thức thu gọn: nhờ AI đổi sang SMILES
            m, err = parse_student(smi)
            if m is None:
                return None, None, None, "Không hiểu chất đã nhập. Hãy gõ công thức phân tử (ví dụ C4H10O), SMILES hoặc tên chất."
    c, h = mol_counts(m)
    if any(e not in VALENCE for e in c):
        return None, None, None, "Chất có nguyên tố chưa được hỗ trợ."
    return c, h, Chem.MolToSmiles(m, isomericSmiles=False), None


def ai_names(smiles_list):
    """Nhờ AI đặt tên. Cấu trúc do thuật toán sinh nên chỉ phần tên có thể sai."""
    lines = "\n".join(f"{i + 1}. {s_}" for i, s_ in enumerate(smiles_list))
    r = ask_ai([RULES_ISO + "\n\nĐặt tên theo danh pháp SGK cho từng chất sau (theo đúng thứ tự). "
                'Chỉ trả JSON {"ten":["tên 1","tên 2",...]} có đúng ' + str(len(smiles_list)) + " phần tử.\n" + lines], temperature=0)
    names = r.get("ten") or []
    return {s_: (names[i] if i < len(names) else "") for i, s_ in enumerate(smiles_list)}


_SMI_RE = re.compile(r"[A-Za-z0-9@+\-\[\]()=#$/\\%.:]+")


def smiles_like(t):
    """Chuỗi giống SMILES thật: không có khoảng trắng hoặc chữ có dấu. Tránh câu tiếng Việt như 'Có bao nhiêu...' bị RDKit đọc nhầm."""
    t = (t or "").strip()
    return bool(t) and _SMI_RE.fullmatch(t) is not None and Chem.MolFromSmiles(t) is not None


def to_smiles(img, text):
    if smiles_like(text):
        return text
    prompt = RULES + '\n\nChuyển chất sau thành SMILES chuẩn. Chỉ trả JSON {"smiles":"","ten":""}, để rỗng nếu không đọc được.'
    if img:
        r = ask_ai([types.Part.from_bytes(data=to_jpeg(img), mime_type="image/jpeg"), prompt])
    elif text:
        r = ask_ai([prompt + "\nChất: " + text])
    else:
        return ""
    return (r.get("smiles") or "").strip()


REL = {
    "cung_chat": ("warning", "Hai công thức biểu diễn **cùng một chất**, không phải đồng phân."),
    "khac_cong_thuc": ("error", "**Không phải đồng phân** vì khác công thức phân tử."),
    "cau_tao": ("success", "Hai chất là **đồng phân cấu tạo** của nhau."),
    "hinh_hoc": ("success", "Hai chất là **đồng phân hình học (cis–trans)** của nhau."),
    "doi_quang": ("success", "Hai chất là **đồng phân quang học (đối quang)** của nhau."),
    "lap_the": ("success", "Hai chất là **đồng phân lập thể** (không phải đối quang) của nhau."),
}


def compare_ui_result(s1, s2):
    m1, m2 = Chem.MolFromSmiles(s1), Chem.MolFromSmiles(s2)
    if m1 is None or m2 is None:
        st.error("Không đọc được một trong hai chất. Hãy kiểm tra lại ảnh hoặc nhập SMILES.")
        return
    c1, c2 = st.columns(2)
    for col, m, s_, tag in ((c1, m1, s1, "A"), (c2, m2, s2, "B")):
        with col:
            st.markdown(f"**Chất {tag}**")
            st.image(Draw.MolToImage(m, size=(320, 240)))
            st.markdown(f"{pretty(formula_of(m))} · M = {Descriptors.MolWt(m):.2f} g/mol")
            st.code(s_, language=None)
    rel = relation(m1, m2)
    kind, msg = REL[rel]
    getattr(st, kind)(msg)
    if rel == "khac_cong_thuc":
        st.markdown(f"{pretty(formula_of(m1))} ≠ {pretty(formula_of(m2))}")
        return
    if rel == "cung_chat":
        return
    try:
        r = ask_ai([RULES_ISO + f'\n\nChất A: {s1}\nChất B: {s2}\nRDKit đã xác nhận quan hệ: {msg}\n'
                    'Hãy đặt tên hai chất và giải thích ngắn gọn theo SGK. Chỉ trả JSON '
                    '{"ten_A":"","ten_B":"","loai_dong_phan":"ví dụ: đồng phân nhóm chức","giai_thich":"2-3 câu: điểm khác nhau về cấu tạo hoặc không gian"}'])
        st.markdown(f"**Tên:** A = {r.get('ten_A', '')}; B = {r.get('ten_B', '')}")
        if r.get("loai_dong_phan"):
            st.markdown("**Loại đồng phân:** " + r["loai_dong_phan"])
        if r.get("giai_thich"):
            st.markdown(r["giai_thich"])
    except Exception as e:
        st.caption(f"Không lấy được phần giải thích: {type(e).__name__}")


def build_iso_result(fkey, orig, p):
    """Bộ máy hóa học: sinh toàn bộ đồng phân, kiểm tra và đếm theo tham số p. AI không tham gia."""
    out = iso_enumerate(fkey, p["ben"])
    if not out["ok"]:
        return {"ok": False, "why": out["why"], "fkey": fkey}
    kept, excluded, hidden = [], {}, 0
    for smi in out["smiles"]:
        m = Chem.MolFromSmiles(smi)
        if p["ben"]:
            bad = next((n for n, pt in _UNST if m.HasSubstructMatch(pt)), None)
            if bad:
                excluded[bad] = excluded.get(bad, 0) + 1
                continue
        labels = iso_labels(m)
        if not pass_filters(m, labels, p):
            hidden += 1
            continue
        tot, geo, opt = stereo_split(m)
        kept.append({"smiles": smi, "labels": labels, "stereo": tot, "geo": geo, "opt": opt})
    return {"ok": True, "fkey": fkey, "params": p, "kept": kept, "excluded": excluded, "hidden": hidden,
            "total_math": len(out["smiles"]), "orig": orig}


LB_CHUC = {"don": "đơn chức", "da": "đa chức hoặc tạp chức", "bat_ky": "mọi loại chức"}
LB_MACH = {"ho": "mạch hở", "vong": "mạch vòng", "ho_hoac_thom": "mạch hở hoặc vòng benzen", "bat_ky": "mọi loại mạch"}
LB_LT = {"khong": "không tính đồng phân lập thể", "hinh_hoc": "kể cả đồng phân hình học (cis–trans)",
         "quang_hoc": "kể cả đồng phân quang học", "tat_ca": "kể cả mọi đồng phân lập thể"}


def answer_of(r):
    k, lt = r["kept"], r["params"]["lap_the"]
    return {"hinh_hoc": sum(x["geo"] for x in k), "quang_hoc": sum(x["opt"] for x in k),
            "tat_ca": sum(x["stereo"] for x in k)}.get(lt, len(k))


def render_iso(r, names):
    kept, p = r["kept"], r["params"]
    nh = " hoặc ".join(p["nhom_ten"])
    st.success(f"**{answer_of(r)} đồng phân** của {pretty(r['fkey'])}" + (f" ({nh})" if nh else ""))
    hid = r["hidden"] + (sum(r["excluded"].values()) if p["ben"] else 0)
    st.caption(f"Cấu tạo {len(kept)} · kể cả hình học {sum(k['geo'] for k in kept)} · quang học {sum(k['opt'] for k in kept)} · "
               f"lập thể {sum(k['stereo'] for k in kept)} · đã ẩn hoặc loại {hid}")
    comp = {}
    for k in kept:
        for lb in k["labels"]:
            d = comp.setdefault(lb, [0, 0])
            d[0] += 1
            d[1] += k["stereo"]
    with st.expander("Thống kê theo nhóm"):
        st.dataframe([{"Nhóm": n, "Cấu tạo": v[0], "Kể cả lập thể": v[1]} for n, v in sorted(comp.items())], hide_index=True, width="stretch")
        if r["params"]["ben"] and r["excluded"]:
            st.caption("Đã loại cấu trúc kém bền: " + ", ".join(f"{n} ({v})" for n, v in r["excluded"].items()))
    pick = st.multiselect("Lọc theo nhóm", sorted(comp), key="iso_filter")
    shown = [k for k in kept if not pick or any(lb in pick for lb in k["labels"])]
    if pick:
        st.caption(f"{len(shown)} đồng phân cấu tạo thuộc nhóm đã chọn.")
    orig = r.get("orig")
    hit = next((i for i, k in enumerate(kept, 1) if k["smiles"] == orig), None) if orig else None
    if hit:
        st.info(f"Chất bạn nhập là đồng phân số {hit}.")
    rows = [{"STT": i, "Nhóm": " + ".join(k["labels"]),
             "Lập thể": "không" if k["stereo"] == 1 else f"{k['stereo']} dạng",
             **({"Tên (AI)": names.get(k["smiles"], "")} if names else {}), "SMILES": k["smiles"]} for i, k in enumerate(shown, 1)]
    st.dataframe(rows[:200], hide_index=True, width="stretch")
    if len(rows) > 200:
        st.caption(f"Hiện 200 dòng đầu trong {len(rows)}. Hãy lọc theo nhóm.")
    if rows:
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
        st.download_button("⬇️ Tải CSV", buf.getvalue().encode("utf-8-sig"), file_name=f"dong_phan_{r['fkey']}.csv", mime="text/csv", key="iso_dl")
    ms = [Chem.MolFromSmiles(k["smiles"]) for k in shown[:48]]
    if ms:
        try:
            st.image(Draw.MolsToGridImage(ms, molsPerRow=4, subImgSize=(230, 180), legends=[str(i) for i in range(1, len(ms) + 1)]))
        except Exception:
            pass


def render_understood(u, r):
    """Một dòng cho biết máy hiểu câu hỏi thế nào (chỉ khi AI đọc câu hỏi)."""
    if not u:
        return
    p = r["params"]
    dflt = set(p.get("defaults", []))
    nhom = (" và " if p["kieu"] == "tat_ca" else " hoặc ").join(p["nhom_ten"]) or "mọi nhóm"
    bits = [f"nhóm: {nhom}", LB_CHUC[p["chuc"]] + ("*" if "chuc" in dflt else ""), LB_MACH[p["mach"]] + ("*" if "mach" in dflt else ""),
            LB_LT[p["lap_the"]] + ("*" if "lap_the" in dflt else "")]
    st.caption("Máy hiểu: " + " · ".join(bits) + ("  (*: mặc định vì câu hỏi không nêu)" if dflt else ""))
    if u.get("chat_goc") and r.get("orig"):
        st.caption(f"Chất được nhắc tới: {u['chat_goc']}")
        try:
            st.image(Draw.MolToImage(Chem.MolFromSmiles(r["orig"]), size=(220, 150)))
        except Exception:
            pass
    for w in u.get("warn") or []:
        st.caption(w)


def isomer_ui():
    top = st.container()  # lỗi và lý do sai hiện ở đầu tab
    t1, t2 = st.tabs(["Liệt kê đồng phân", "So sánh hai chất"])
    with t1:
        q = st.text_input("Câu hỏi hoặc công thức", key="iso_q",
                          placeholder="Ví dụ: số đồng phân este của C4H8O2   ·   C4H10O").strip()
        with st.expander("Tùy chọn"):
            scope = st.radio("Phạm vi khi chỉ nhập công thức", ["Đơn chức, mạch hở (đề SGK)", "Phổ thông (cả đa chức và 1 vòng)",
                                                               "Toán học (mọi cấu trúc)"], key="iso_scope")
            want_names = st.checkbox("Gọi tên bằng AI (có thể sai)", key="iso_names")
            f = st.file_uploader("Hoặc dùng ảnh công thức", type=["jpg", "jpeg", "png", "webp"], key="iso_up")
        img = f.getvalue() if f else None
        if st.button("Đếm đồng phân", type="primary", disabled=not (q or img), key="iso_go"):
            direct = bool(img) or is_simple_formula(q) or smiles_like(q)
            ai_used = bool(img or want_names or not direct)
            if ai_used and st.session_state.n >= MAX_PER_SESSION:
                show_top_error(top, "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục.")
            else:
                with st.spinner("Đang tính…"):
                    try:
                        st.session_state.iso = None
                        counts = nH = orig = err = und = None
                        if direct:
                            counts, nH, orig, err = resolve_formula(img, q)
                            p = dict(PRESETS["mono" if scope.startswith("Đơn") else ("pho" if scope.startswith("Phổ") else "math")])
                        else:
                            p, warn, raw = parse_question(q)
                            ftxt, gc = raw.get("cong_thuc"), raw.get("chat_goc")
                            und = {"warn": warn, "chat_goc": gc}
                            if ftxt and is_simple_formula(str(ftxt)):
                                counts, nH = parse_formula(str(ftxt))
                            elif gc:
                                counts, nH, orig, err = resolve_formula(None, str(gc))
                            else:
                                err = raw.get("can_hoi_lai") or "Câu hỏi chưa có công thức hoặc tên chất."
                        if err:
                            show_top_error(top, err)
                        else:
                            n_heavy = sum(counts.values())
                            fkey = hill_key(counts, nH)
                            if n_heavy > 12:
                                show_top_error(top, f"{pretty(fkey)} có {n_heavy} nguyên tử khác H, quá lớn để liệt kê chính xác (tối đa 12).")
                            else:
                                res = build_iso_result(fkey, orig, p)
                                if not res["ok"]:
                                    show_top_error(top, f"{pretty(fkey)} có quá nhiều đồng phân để sinh trong {ISO_TIME:.0f} giây. Hãy nêu rõ nhóm chức.")
                                else:
                                    nm = {}
                                    if want_names and 0 < len(res["kept"]) <= 60:
                                        nm = ai_names([k["smiles"] for k in res["kept"]])
                                    st.session_state.iso, st.session_state.iso_nm, st.session_state.iso_und = res, nm, und
                                    add_history(pretty(res['fkey']), 'Đồng phân')
                                    if ai_used:
                                        st.session_state.n += 1
                    except Exception as e:
                        st.session_state.iso = None
                        show_top_error(top, "Có lỗi, hãy thử lại.", e)
        res = st.session_state.get("iso")
        if res:
            st.divider()
            render_understood(st.session_state.get("iso_und"), res)
            if res["kept"]:
                render_iso(res, st.session_state.get("iso_nm") or {})
            else:
                st.warning(f"Không có đồng phân nào của {pretty(res['fkey'])} thỏa điều kiện (đã ẩn {res['hidden']}).")
    with t2:
        st.caption("Mỗi chất nhập bằng ảnh, SMILES hoặc tên.")
        ca, cb = st.columns(2)
        inp = []
        for col, tag in ((ca, "A"), (cb, "B")):
            with col:
                st.markdown(f"**Chất {tag}**")
                f = st.file_uploader(f"Ảnh chất {tag}", type=["jpg", "jpeg", "png", "webp"], key=f"cmp_up{tag}")
                t = st.text_input(f"Hoặc SMILES/tên chất {tag}", key=f"cmp_t{tag}").strip()
                inp.append((f.getvalue() if f else None, t))
        ok = all(i or t for i, t in inp)
        if st.button("So sánh", type="primary", disabled=not ok, key="cmp_go"):
            if st.session_state.n >= MAX_PER_SESSION:
                show_top_error(top, "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục.")
            else:
                with st.spinner("Đang phân tích…"):
                    try:
                        smi = [to_smiles(i, t) for i, t in inp]
                        st.session_state.n += 1
                        st.session_state.cmp = smi
                    except Exception as e:
                        st.session_state.cmp = None
                        show_top_error(top, "Có lỗi khi đọc chất, hãy thử lại.", e)
        if st.session_state.get("cmp"):
            st.divider()
            compare_ui_result(*st.session_state.cmp)



# ---------------- Dụng cụ thí nghiệm ----------------
RULES_EQ = """Bạn là giáo viên Hóa học Việt Nam hướng dẫn thực hành thí nghiệm ở trường phổ thông. Hãy nhận biết dụng cụ thí nghiệm theo tên gọi trong SGK Việt Nam, ví dụ: ống nghiệm, cốc thủy tinh, bình tam giác, bình cầu (đáy tròn, đáy bằng, có nhánh), phễu lọc, phễu chiết, đèn cồn, kẹp ống nghiệm, giá sắt, kiềng đun, lưới amiăng, ống đong, pipet, buret, bình định mức, cối và chày sứ, chén sứ, đũa thủy tinh, thìa lấy hóa chất, nhiệt kế, ống sinh hàn, bình Kíp, ống dẫn khí, nút cao su, giấy lọc, cân, bình tia, giá để ống nghiệm. Chỉ nêu những gì thấy rõ trong ảnh; nếu không chắc thì giảm độ tin cậy và nói rõ."""

EFMT = """Chỉ trả về MỘT đối tượng JSON, không markdown, không lời dẫn:
{"loai_anh":"dụng cụ riêng lẻ | bộ lắp thí nghiệm | không phải dụng cụ thí nghiệm",
"dung_cu":[{"ten":"tên theo SGK","ten_khac":"tên gọi khác hoặc tên tiếng Anh, có thể rỗng","vi_tri":"vị trí trong ảnh, ví dụ: bên trái, ở giữa","chat_lieu":"thủy tinh | sứ | kim loại | nhựa | cao su | ...","cong_dung":"1-2 câu","cach_dung":["bước hoặc lưu ý sử dụng"],"luu_y_an_toan":["điều cần chú ý để an toàn"],"do_tin_cay":số 0-1}],
"thi_nghiem":null hoặc {"ten":"thí nghiệm có thể đang lắp, ví dụ: điều chế khí oxi từ KMnO4","mo_ta":"1-2 câu","nhan_xet_lap_rap":["điểm đúng hoặc điểm có thể sai/nguy hiểm trong cách lắp, chỉ nêu khi nhìn thấy rõ"]},
"luu_y":["điều chưa chắc chắn, ví dụ ảnh mờ, bị che khuất"]}
Nếu ảnh có nhiều dụng cụ thì liệt kê từng dụng cụ nhận ra được. Chỉ đặt thi_nghiem khi ảnh là một bộ lắp ráp rõ ràng. Nếu không phải dụng cụ thí nghiệm, để dung_cu rỗng và giải thích trong luu_y."""


def analyze_equipment(img, text):
    head = RULES_EQ + "\n\n"
    if img:
        return ask_ai([types.Part.from_bytes(data=to_jpeg(img), mime_type="image/jpeg"),
                       head + "Đây là ảnh dụng cụ thí nghiệm hóa học. Hãy nhận biết.\n" + EFMT])
    return ask_ai([head + "Hãy giải thích dụng cụ có tên: " + text + ".\n" + EFMT])


def render_equipment(r):
    kind = r.get("loai_anh")
    tools = r.get("dung_cu") or []
    if kind:
        st.markdown(f"**Loại ảnh:** {kind}")
    if not tools:
        st.warning("Chưa nhận ra dụng cụ thí nghiệm nào trong ảnh.")
    for t in tools:
        with st.container(border=True):
            st.markdown(f"### {t.get('ten', 'Dụng cụ')}")
            if t.get("ten_khac"):
                st.caption(f"Tên khác: {t['ten_khac']}")
            meta = [x for x in (t.get("vi_tri"), t.get("chat_lieu")) if x]
            if meta:
                st.markdown("**Vị trí / chất liệu:** " + " · ".join(meta))
            conf = t.get("do_tin_cay")
            if isinstance(conf, (int, float)):
                p = max(0.0, min(1.0, float(conf)))
                st.progress(p, text=f"Độ tin cậy: {round(p * 100)}%")
                if p < 0.6:
                    st.warning("Độ tin cậy thấp, hãy hỏi giáo viên hoặc chụp lại rõ hơn.")
            if t.get("cong_dung"):
                st.markdown("**Công dụng:** " + t["cong_dung"])
            bullets("Cách dùng", t.get("cach_dung"))
            if t.get("luu_y_an_toan"):
                st.warning("**An toàn:** " + " • ".join([t["luu_y_an_toan"]] if isinstance(t["luu_y_an_toan"], str) else t["luu_y_an_toan"]))
    tn = r.get("thi_nghiem")
    if isinstance(tn, dict) and (tn.get("ten") or tn.get("mo_ta")):
        st.divider()
        st.markdown("### Bộ thí nghiệm trong ảnh")
        if tn.get("ten"):
            st.markdown("**Có thể là:** " + tn["ten"])
        if tn.get("mo_ta"):
            st.markdown(tn["mo_ta"])
        bullets("Nhận xét về cách lắp", tn.get("nhan_xet_lap_rap"))
    if r.get("luu_y"):
        st.info("Lưu ý: " + " • ".join([r["luu_y"]] if isinstance(r["luu_y"], str) else r["luu_y"]))
    st.caption("Kết quả do AI nhận biết qua ảnh nên có thể nhầm, nhất là khi ảnh mờ hoặc dụng cụ bị che. "
               "Khi làm thí nghiệm thật, hãy làm theo hướng dẫn của giáo viên và quy định an toàn của phòng thực hành.")


def equipment_ui():
    top = st.container()  # lỗi và lý do sai hiện ở đầu tab
    src = st.radio("Nguồn", ["🖼️ Tải ảnh", "📷 Chụp ảnh", "⌨️ Gõ tên dụng cụ"], horizontal=True, key="eq_src")
    img, txt = None, None
    if src.startswith("📷"):
        f = st.camera_input("Chụp dụng cụ hoặc bộ thí nghiệm", key="eq_cam")
        img = f.getvalue() if f else None
    elif src.startswith("🖼️"):
        f = st.file_uploader("Ảnh dụng cụ hoặc bộ thí nghiệm", type=["jpg", "jpeg", "png", "webp"], key="eq_up")
        img = f.getvalue() if f else None
    else:
        txt = st.text_input("Tên dụng cụ", placeholder="Ví dụ: ống sinh hàn", key="eq_txt").strip()
    if st.button("Nhận biết dụng cụ", type="primary", disabled=not (img or txt), key="eq_go"):
        if st.session_state.n >= MAX_PER_SESSION:
            show_top_error(top, "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục.")
        else:
            with st.spinner("Đang phân tích…"):
                try:
                    st.session_state.eq = analyze_equipment(img, txt)
                    add_history(txt or ', '.join(str(t.get('ten', '')) for t in ((st.session_state.eq or {}).get('dung_cu') or [])[:2]) or 'Dụng cụ', 'Dụng cụ')
                    st.session_state.n += 1
                except Exception as e:
                    st.session_state.eq = None
                    show_top_error(top, "Có lỗi khi phân tích, hãy thử lại hoặc chụp rõ hơn.", e)
    if st.session_state.get("eq"):
        st.divider()
        render_equipment(st.session_state.eq)



# ---------------- Chẩn đoán lỗi & chấm bài ----------------
GROUPS = [
    ("ancol", "[CX4][OX2H]"), ("phenol", "c[OX2H]"), ("ete", "[CX4][OX2][CX4]"),
    ("anđehit", "[CX3;H1,H2;!$(C[OX2,NX3])]=O"), ("xeton", "[#6][CX3](=O)[#6]"),
    ("axit cacboxylic", "[CX3](=O)[OX2H1]"), ("este", "[CX3](=O)[OX2][#6]"),
    ("amin", "[NX3;!$(NC=[O,N,S])][#6]"), ("nitro", "[N+](=O)[O-]"),
    ("anken (C=C)", "[CX3]=[CX3]"), ("ankin (C≡C)", "[CX2]#[CX2]"),
    ("vòng benzen", "c1ccccc1"), ("dẫn xuất halogen", "[#6][F,Cl,Br,I]"),
]
_GP = [(n, Chem.MolFromSmarts(p)) for n, p in GROUPS]


def groups_of(m):
    return [n for n, p in _GP if m.HasSubstructMatch(p)]


def atom_counts(m):
    c = {}
    for a in Chem.AddHs(m).GetAtoms():
        c[a.GetSymbol()] = c.get(a.GetSymbol(), 0) + 1
    return c


def diff_text(ms, mt):
    a, b = atom_counts(ms), atom_counts(mt)
    order = ["C", "H"] + sorted(k for k in set(a) | set(b) if k not in ("C", "H"))
    out = []
    for k in order:
        d = a.get(k, 0) - b.get(k, 0)
        if d > 0:
            out.append(f"thừa {d} nguyên tử {k}")
        elif d < 0:
            out.append(f"thiếu {-d} nguyên tử {k}")
    return ", ".join(out)


def parse_student(smi):
    """Trả về (mol, thông báo lỗi bằng tiếng Việt)."""
    smi = (smi or "").strip()
    if not smi:
        return None, "Chưa đọc được công thức em vẽ."
    m = Chem.MolFromSmiles(smi, sanitize=False)
    if m is None:
        return None, "Công thức bị sai cú pháp (thiếu hoặc thừa ngoặc, vòng chưa đóng hoặc ký hiệu lạ)."
    try:
        Chem.SanitizeMol(m)
    except Chem.AtomValenceException as e:
        mm = re.search(r"atom # (\d+) (\w+), (\d+)", str(e))
        if mm:
            return None, f"Nguyên tử {mm.group(2)} thứ {int(mm.group(1)) + 1} có hóa trị {mm.group(3)}, vượt quá hóa trị cho phép."
        return None, "Có nguyên tử vượt quá hóa trị cho phép."
    except Chem.KekulizeException:
        return None, "Vòng thơm vẽ sai: không sắp xếp được các nối đôi xen kẽ trong vòng."
    except Exception:
        return None, "Công thức có lỗi hóa trị."
    return m, None


def draw_student(smi, size=(320, 220)):
    """Vẽ lại đúng hình học sinh đã vẽ, kể cả khi sai hóa trị; khoanh đỏ nguyên tử sai nếu có."""
    m = Chem.MolFromSmiles((smi or "").strip(), sanitize=False)
    if m is None:
        return None
    bad = []
    try:
        Chem.SanitizeMol(m)
    except Chem.AtomValenceException as e:
        mm = re.search(r"atom # (\d+)", str(e))
        if mm:
            bad = [int(mm.group(1))]
    except Exception:
        pass
    try:
        m.UpdatePropertyCache(strict=False)
        Chem.FastFindRings(m)
        return Draw.MolToImage(m, size=size, highlightAtoms=bad, highlightColor=(1.0, 0.55, 0.55))
    except Exception:
        return None


def grade_structure(smi, target_smi):
    mt = Chem.MolFromSmiles(target_smi)
    ms, err = parse_student(smi)
    tips = ["Nhớ hóa trị: C có 4 liên kết, N có 3, O có 2, H và halogen có 1."]
    if ms is None:
        return {"kq": "loi", "nx": [err], "goi_y": tips}
    fs, ft = formula_of(ms), formula_of(mt)
    if noiso(ms) == noiso(mt):
        return {"kq": "dung", "nx": ["Chính xác. Công thức cấu tạo em vẽ đúng."], "goi_y": []}
    gs, gt = set(groups_of(ms)), set(groups_of(mt))
    if fs == ft:
        nx = [f"Công thức phân tử đúng ({pretty(fs)}) nhưng cấu tạo chưa đúng yêu cầu: em đã vẽ một đồng phân khác."]
        if gs != gt:
            nx.append("Nhóm chức khác với chất cần tìm. Hình em vẽ có: " + (", ".join(sorted(gs)) or "không có nhóm chức đặc trưng")
                      + ". Đề cần: " + (", ".join(sorted(gt)) or "không có nhóm chức đặc trưng") + ".")
            nx.append("Loại đồng phân em đã vẽ: đồng phân nhóm chức.")
        else:
            nx.append("Nhóm chức đúng nhưng mạch cacbon hoặc vị trí nhóm thế/liên kết bội chưa đúng.")
        return {"kq": "dong_phan", "nx": nx, "goi_y": ["Đếm lại số cacbon ở mạch chính và vị trí gắn nhóm thế."]}
    nx = [f"Công thức phân tử của hình em vẽ là {pretty(fs)}. So với chất cần tìm, em {diff_text(ms, mt)}."]
    if gs != gt:
        nx.append("Nhóm chức chưa khớp: hình em vẽ có " + (", ".join(sorted(gs)) or "không có nhóm chức đặc trưng")
                  + "; đề cần " + (", ".join(sorted(gt)) or "không có nhóm chức đặc trưng") + ".")
    return {"kq": "sai_cong_thuc", "nx": nx, "goi_y": ["Đếm lại số nguyên tử C và số H gắn vào mỗi C."]}


def norm_name(x):
    x = unicodedata.normalize("NFD", (x or "").lower().replace("đ", "d"))
    x = "".join(c for c in x if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]", "", x)


BANK = [
    ("metan", "C", []), ("etan", "CC", []), ("propan", "CCC", []), ("butan", "CCCC", ["n-butan"]),
    ("2-metylpropan", "CC(C)C", ["isobutan"]), ("pentan", "CCCCC", ["n-pentan"]),
    ("2-metylbutan", "CCC(C)C", ["isopentan"]), ("2,2-đimetylpropan", "CC(C)(C)C", ["neopentan"]),
    ("eten", "C=C", ["etilen"]), ("propen", "C=CC", ["propilen"]), ("but-1-en", "C=CCC", ["1-buten"]),
    ("but-2-en", "CC=CC", ["2-buten"]), ("etin", "C#C", ["axetilen"]), ("propin", "C#CC", []),
    ("benzen", "c1ccccc1", []), ("toluen", "Cc1ccccc1", ["metylbenzen"]), ("stiren", "C=Cc1ccccc1", ["vinylbenzen"]),
    ("phenol", "Oc1ccccc1", ["hiđroxibenzen"]), ("anilin", "Nc1ccccc1", ["phenylamin"]),
    ("metanol", "CO", ["ancol metylic"]), ("etanol", "CCO", ["ancol etylic"]),
    ("propan-1-ol", "CCCO", ["ancol propylic"]), ("propan-2-ol", "CC(C)O", ["ancol isopropylic"]),
    ("butan-1-ol", "CCCCO", []), ("etylen glicol", "OCCO", ["etan-1,2-điol"]),
    ("glixerol", "OCC(O)CO", ["glixerin", "propan-1,2,3-triol"]),
    ("metanal", "C=O", ["fomanđehit", "anđehit fomic"]), ("etanal", "CC=O", ["anđehit axetic", "axetanđehit"]),
    ("propanon", "CC(C)=O", ["axeton"]), ("axit metanoic", "OC=O", ["axit fomic"]),
    ("axit etanoic", "CC(=O)O", ["axit axetic"]), ("axit propanoic", "CCC(=O)O", ["axit propionic"]),
    ("axit benzoic", "OC(=O)c1ccccc1", []), ("metyl axetat", "COC(C)=O", ["metyl etanoat"]),
    ("etyl axetat", "CCOC(C)=O", ["etyl etanoat"]), ("metylamin", "CN", ["metanamin"]), ("etylamin", "CCN", ["etanamin"]),
    ("glyxin", "NCC(=O)O", ["axit aminoaxetic"]), ("alanin", "CC(N)C(=O)O", ["axit 2-aminopropanoic"]),
    ("đimetyl ete", "COC", []), ("đietyl ete", "CCOCC", []), ("etyl clorua", "CCCl", ["cloetan"]),
    ("clorofom", "ClC(Cl)Cl", ["triclometan"]),
]

GRAPH_PROMPT = RULES + """

Đây là ảnh BÀI LÀM của học sinh và sẽ được chấm điểm. Hãy CHÉP LẠI TRUNG THỰC cấu trúc học sinh đã vẽ dưới dạng ĐỒ THỊ nguyên tử - liên kết. TUYỆT ĐỐI không sửa bài, không đoán học sinh định vẽ chất nào và không tự cân lại hóa trị. Nếu hình vẽ sai (ví dụ nguyên tử C nối với 5 nét) thì vẫn chép đúng 5 nét đó.

Quy tắc:
- Mỗi nguyên tử trong hình là một đối tượng {"id":số thứ tự bắt đầu từ 1,"el":"C" hoặc "O", "N", "Cl", "Br"...,"h":số H}.
- Công thức khung (đường gấp khúc): mỗi đầu mút và mỗi chỗ gấp khúc là một nguyên tử C với "h":null (H ngầm định, không vẽ).
- Công thức thu gọn hoặc khai triển: "h" là SỐ H được ghi kèm hoặc vẽ nối trực tiếp vào nguyên tử đó. Ví dụ CH3 thì h=3; nếu có 4 nét nối từ C tới 4 chữ H thì h=4. Đếm đúng từng nét, không làm tròn theo hóa trị. Nguyên tử không ghi và không vẽ H nào thì "h":null.
- KHÔNG tạo đối tượng riêng cho các H đã tính vào "h".
- Mỗi liên kết giữa hai nguyên tử (không kể H) là [id1,id2,bậc], bậc 1, 2 hoặc 3 tính theo số nét song song.
- Vòng benzen vẽ có vòng tròn bên trong: ghi 6 nguyên tử C và các liên kết xen kẽ bậc 1 và 2.
Chỉ trả về JSON: {"atoms":[...],"bonds":[[1,2,1],...],"ghi_chu":"điều chưa rõ hoặc mơ hồ trong hình, có thể rỗng"}"""


def graph_to_smiles(g):
    """Dựng phân tử từ đồ thị AI chép lại, giữ nguyên số H đã vẽ để lỗi hóa trị không bị che mất."""
    atoms, bonds = g.get("atoms") or [], g.get("bonds") or []
    if not atoms:
        return ""
    rw, idx = Chem.RWMol(), {}
    for a in atoms:
        el = str(a.get("el", "C")).strip()
        at = Chem.Atom(el[:1].upper() + el[1:].lower())
        h = a.get("h")
        if isinstance(h, (int, float)) and not isinstance(h, bool):
            at.SetNumExplicitHs(int(h))
            at.SetNoImplicit(True)
        idx[a["id"]] = rw.AddAtom(at)
    kinds = {1: Chem.BondType.SINGLE, 2: Chem.BondType.DOUBLE, 3: Chem.BondType.TRIPLE}
    for b in bonds:
        i, j, o = b[0], b[1], int(b[2])
        if i != j and rw.GetBondBetweenAtoms(idx[i], idx[j]) is None:
            rw.AddBond(idx[i], idx[j], kinds[o])
    m = rw.GetMol()
    m.UpdatePropertyCache(strict=False)
    Chem.FastFindRings(m)
    return Chem.MolToSmiles(m)


def transcribe_student(img):
    r = ask_ai([types.Part.from_bytes(data=to_jpeg(img), mime_type="image/jpeg"), GRAPH_PROMPT], temperature=0)
    try:
        smi = graph_to_smiles(r)
    except Exception:
        raise ValueError("Không dựng được cấu trúc từ ảnh, hãy chụp rõ hơn.")
    return smi, r.get("ghi_chu", "")


def analyze_image_strict(img):
    """Đọc ảnh trung thực (đồ thị nguyên tử) rồi mới gọi tên và tính chất nếu công thức hợp lệ."""
    smi, note = transcribe_student(img)
    res = {"smiles": smi}
    if smi and parse_student(smi)[0] is not None:
        res = analyze_smiles(smi)
        res["smiles"] = smi
        res.pop("do_tin_cay", None)
    if note:
        old = res.get("luu_y") or []
        res["luu_y"] = ([old] if isinstance(old, str) else list(old)) + [note]
    return res


RANDOM_Q = "🎲 Ngẫu nhiên"


def new_question(kind):
    st.session_state.pq = {"kind": kind, "idx": random.randrange(len(BANK)), "result": None}
    st.session_state["pr_pick"] = RANDOM_Q


def practice_ui():
    st.session_state.setdefault("pq", None)
    st.session_state.setdefault("ph", [])
    top = st.container()  # kết quả chấm và lý do sai hiện ở đầu tab
    c1, c2 = st.columns([3, 2])
    with c1:
        kind = st.radio("Dạng bài", ["✍️ Vẽ công thức cấu tạo", "🏷️ Gọi tên chất"], horizontal=True, key="pr_kind")
    with c2:
        who = st.text_input("Tên học sinh (không bắt buộc)", key="pr_name")
    if st.session_state.pq is None or st.session_state.pq["kind"] != kind:
        new_question(kind)
    names = [b[0] for b in BANK]
    pick = st.selectbox("Chọn đề cụ thể (hoặc để ngẫu nhiên)", [RANDOM_Q] + names, key="pr_pick")
    if pick != RANDOM_Q and names.index(pick) != st.session_state.pq["idx"]:
        st.session_state.pq = {"kind": kind, "idx": names.index(pick), "result": None}
    q = st.session_state.pq
    ten, smi, alias = BANK[q["idx"]]
    mt = Chem.MolFromSmiles(smi)

    with st.container(border=True):
        if kind.startswith("✍️"):
            st.markdown(f"#### Viết công thức cấu tạo của: **{ten}**")
            if st.checkbox("Gợi ý: hiện công thức phân tử", key=f"pr_hint_{q['idx']}"):
                st.markdown(f"Công thức phân tử: **{pretty(formula_of(mt))}**")
            how = st.radio("Nộp bài bằng", ["🖼️ Ảnh bài làm", "📷 Chụp ảnh", "⌨️ Nhập SMILES"], horizontal=True, key="pr_how")
            img, ans = None, ""
            if how.startswith("🖼️"):
                f = st.file_uploader("Ảnh bài làm", type=["jpg", "jpeg", "png", "webp"], key="pr_up")
                img = f.getvalue() if f else None
            elif how.startswith("📷"):
                f = st.camera_input("Chụp bài làm", key="pr_cam")
                img = f.getvalue() if f else None
            else:
                ans = st.text_input("SMILES bài làm", key="pr_smi").strip()
            go = st.button("Nộp bài", type="primary", disabled=not (img or ans), key="pr_go")
        else:
            st.markdown("#### Gọi tên chất sau")
            st.image(Draw.MolToImage(mt, size=(360, 240)))
            ans = st.text_input("Tên chất (theo SGK)", key="pr_name_ans").strip()
            img = None
            go = st.button("Nộp bài", type="primary", disabled=not ans, key="pr_go2")

    if go and q["result"] is None:
        if img and st.session_state.n >= MAX_PER_SESSION:
            show_top_error(top, "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục.")
        else:
            try:
                note, given = "", ans
                if kind.startswith("✍️"):
                    if img:
                        with st.spinner("Đang đọc bài làm…"):
                            given, note = transcribe_student(img)
                        st.session_state.n += 1
                    res = grade_structure(given, smi)
                else:
                    ok = norm_name(ans) in {norm_name(x) for x in [ten] + alias}
                    res = {"kq": "dung" if ok else "sai",
                           "nx": ["Chính xác." if ok else "Chưa đúng. Hãy xác định mạch chính, nhóm chức và vị trí nhóm thế rồi gọi tên lại."],
                           "goi_y": []}
                res.update({"given": given, "note": note})
                q["result"] = res
                add_history(ten, 'Luyện tập')
                st.session_state.ph.append({
                    "Thời gian": datetime.now().strftime("%H:%M:%S"), "Học sinh": who, "Dạng bài": "vẽ CTCT" if kind.startswith("✍️") else "gọi tên",
                    "Đề": ten, "Bài làm": given, "Kết quả": res["kq"]})
            except Exception as e:
                show_top_error(top, "Có lỗi khi chấm bài, hãy thử lại.", e)

    with top:
        res = q["result"]
        if res:
            head = {"dung": "Đúng rồi!", "dong_phan": "CHƯA ĐÚNG: đã vẽ một đồng phân khác", "sai_cong_thuc": "CHƯA ĐÚNG: sai công thức phân tử",
                    "loi": "CHƯA ĐÚNG: công thức có lỗi", "sai": "CHƯA ĐÚNG"}[res["kq"]]
            if res["kq"] == "dung":
                st.success("**" + head + "**")
                for t in res["nx"]:
                    st.markdown("- " + t)
            else:
                st.markdown(errbanner(head, *res["nx"]), unsafe_allow_html=True)
            st.markdown("**Đề:** " + (f"viết công thức cấu tạo của **{ten}**" if kind.startswith("✍️") else "gọi tên chất trong hình"))
            for t in res["goi_y"]:
                st.markdown("- 💡 " + t)
            if res.get("note"):
                st.caption("Ghi chú khi đọc ảnh: " + res["note"])
            if kind.startswith("✍️") and res.get("given"):
                im = draw_student(res["given"])
                if im is not None:
                    st.markdown("**Hình em vẽ (máy đọc lại đúng như hình, giữ nguyên lỗi" + (", nguyên tử sai khoanh đỏ" if res["kq"] == "loi" else "") + "):**")
                    st.image(im)
            with st.expander("Xem đáp án"):
                st.image(Draw.MolToImage(mt, size=(300, 200)))
                st.markdown(f"**{ten}** · {pretty(formula_of(mt))} · `{smi}`")
                gl = groups_of(mt)
                if gl:
                    st.markdown("Nhóm chức/đặc điểm: " + ", ".join(gl))
            st.button("➡️ Câu tiếp theo", on_click=new_question, args=(kind,), key="pr_next")

    hist = st.session_state.ph
    if hist:
        st.divider()
        n_ok = sum(1 for h in hist if h["Kết quả"] == "dung")
        a, b, c = st.columns(3)
        a.metric("Đã làm", len(hist))
        b.metric("Đúng", n_ok)
        c.metric("Tỉ lệ đúng", f"{round(100 * n_ok / len(hist))}%")
        st.dataframe(hist, hide_index=True, width="stretch")
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=list(hist[0].keys()))
        w.writeheader()
        w.writerows(hist)
        st.download_button("⬇️ Tải phiếu kết quả (CSV, nộp cho giáo viên)", buf.getvalue().encode("utf-8-sig"),
                           file_name="phieu_ket_qua.csv", mime="text/csv", key="pr_dl")
    st.caption("Phần chấm công thức do RDKit so sánh cấu tạo nên chính xác. Phần đọc ảnh bài làm do AI thực hiện nên có thể đọc sai nét viết khó.")


# ---------------- Đánh giá độ chính xác (giáo viên / nhóm nghiên cứu) ----------------
def parse_truth(text):
    rows = {}
    for line in (text or "").splitlines():
        parts = [p.strip() for p in line.replace(";", ",").split(",")]
        if len(parts) >= 2 and parts[0] and parts[1]:
            rows[parts[0]] = (parts[1], parts[2] if len(parts) > 2 else "")
    return rows


def bench_row(pred, truth):
    mp, mt = Chem.MolFromSmiles(pred or ""), Chem.MolFromSmiles(truth)
    if mp is None or mt is None:
        return {"hop_le": False, "khop_cong_thuc": False, "khop_cau_tao": False, "khop_hoan_toan": False}
    return {"hop_le": True, "khop_cong_thuc": formula_of(mp) == formula_of(mt),
            "khop_cau_tao": noiso(mp) == noiso(mt),
            "khop_hoan_toan": Chem.MolToSmiles(mp) == Chem.MolToSmiles(mt)}


def bench_summary(rows):
    n = len(rows)
    pct = lambda k: round(100 * sum(1 for r in rows if r[k]) / n, 1) if n else 0
    return {"n": n, "hop_le": pct("hop_le"), "khop_cong_thuc": pct("khop_cong_thuc"),
            "khop_cau_tao": pct("khop_cau_tao"), "khop_hoan_toan": pct("khop_hoan_toan"),
            "tb_giay": round(sum(r["giay"] for r in rows) / n, 1) if n else 0}


def bench_ui():
    code = st.secrets.get("TEACHER_CODE", "")
    if code and st.text_input("Mã giáo viên", type="password", key="bm_code") != code:
        st.info("Tab này dành cho giáo viên hoặc nhóm nghiên cứu. Nhập mã giáo viên để dùng.")
        return
    st.markdown("Tải lên một bộ ảnh đã biết đáp án để đo **độ chính xác nhận dạng**. Kết quả có thể đưa vào báo cáo đề tài.")
    files = st.file_uploader("Ảnh kiểm thử", type=["jpg", "jpeg", "png", "webp"], accept_multiple_files=True, key="bm_files")
    up = st.file_uploader("File đáp án CSV (tên_file, SMILES, nhóm)", type=["csv", "txt"], key="bm_csv")
    txt = st.text_area("Hoặc dán đáp án, mỗi dòng: tên_file, SMILES, nhóm", height=130, key="bm_truth",
                       placeholder="03_toluen.jpg, Cc1ccccc1, benzen\n06_etyl_axetat.jpg, CCOC(C)=O, công thức khung")
    method = st.radio("Phương pháp đọc ảnh", ["Đồ thị nguyên tử (trung thực với bài vẽ)", "SMILES trực tiếp"], horizontal=True, key="bm_method")
    src = up.getvalue().decode("utf-8-sig", errors="ignore") if up else txt
    truth = parse_truth(src)
    todo = [f for f in (files or []) if f.name in truth][:BENCH_MAX]
    if files and truth:
        st.caption(f"Có {len(todo)}/{len(files)} ảnh khớp tên trong file đáp án (tối đa {BENCH_MAX} ảnh mỗi lần chạy).")
    if st.button("Chạy đánh giá", type="primary", disabled=not todo, key="bm_go"):
        rows, bar = [], st.progress(0.0, text="Đang chạy…")
        for i, f in enumerate(todo):
            t0 = time.time()
            pred, err = "", ""
            try:
                pred = transcribe_student(f.getvalue())[0] if method.startswith("Đồ thị") else to_smiles(f.getvalue(), "")
            except Exception as e:
                err = type(e).__name__
            tr, grp = truth[f.name]
            r = {"tep": f.name, "nhom": grp, "dap_an": tr, "phuong_phap": "do_thi" if method.startswith("Đồ thị") else "smiles", "ket_qua_may": pred, "loi": err, "giay": round(time.time() - t0, 1)}
            r.update(bench_row(pred, tr))
            rows.append(r)
            bar.progress((i + 1) / len(todo), text=f"Đã xử lý {i + 1}/{len(todo)}")
        st.session_state.bm = rows
    rows = st.session_state.get("bm")
    if rows:
        sm = bench_summary(rows)
        a, b, c, d = st.columns(4)
        a.metric("Số ảnh", sm["n"])
        b.metric("Khớp cấu tạo", f"{sm['khop_cau_tao']}%")
        c.metric("Khớp công thức PT", f"{sm['khop_cong_thuc']}%")
        d.metric("SMILES hợp lệ", f"{sm['hop_le']}%")
        st.caption(f"Khớp hoàn toàn (kể cả lập thể): {sm['khop_hoan_toan']}% · thời gian trung bình: {sm['tb_giay']} giây/ảnh")
        grp = sorted({r["nhom"] for r in rows if r["nhom"]})
        if grp:
            st.markdown("**Theo nhóm nét vẽ**")
            st.dataframe([{"Nhóm": g, **{k: v for k, v in bench_summary([r for r in rows if r["nhom"] == g]).items()
                                         if k in ("n", "khop_cau_tao", "khop_cong_thuc", "hop_le")}} for g in grp],
                         hide_index=True, width="stretch")
        st.dataframe(rows, hide_index=True, width="stretch")
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
        st.download_button("⬇️ Tải bảng kết quả (CSV)", buf.getvalue().encode("utf-8-sig"), file_name="ket_qua_danh_gia.csv",
                           mime="text/csv", key="bm_dl")
    st.caption("Khớp cấu tạo: SMILES máy đọc được và đáp án giống nhau về cấu tạo (RDKit chuẩn hóa trước khi so sánh).")


# ---------------- Sơ đồ tư duy ----------------
MM_COLORS = ["#e11d48", "#f59e0b", "#16a34a", "#0ea5e9", "#8b5cf6", "#ec4899", "#14b8a6", "#f97316"]
MM_FONT = "Arial, Helvetica, sans-serif"

MM_PROMPT = """Bạn là giáo viên Hóa học THPT Việt Nam. Hãy lập SƠ ĐỒ TƯ DUY cho chủ đề được yêu cầu, đúng nội dung SGK Việt Nam.
Chỉ trả về MỘT đối tượng JSON, không markdown:
{"chu_de":"tên chủ đề ngắn gọn","chat_smiles":"SMILES nếu chủ đề là MỘT chất cụ thể (ví dụ etanol), nếu là nhóm chất hoặc khái niệm thì null",
"nhanh":[{"ten":"tên nhánh, tối đa 24 ký tự","y":["ý ngắn, tối đa 80 ký tự"]}]}
Quy tắc:
- Chọn các nhánh phù hợp chủ đề, ví dụ: Khái niệm và cấu tạo, Danh pháp, Tính chất vật lí, Tính chất hóa học, Điều chế, Ứng dụng, Nhận biết, Lưu ý. Với chủ đề phản ứng thì dùng: Định nghĩa, Điều kiện, Ví dụ, Quy tắc, Nhận biết.
- Mỗi ý ngắn gọn, dễ nhớ. Viết công thức bằng chữ thường (CH3-CH2-OH). Phương trình viết dạng: A + B → C (điều kiện).
- KHÔNG nêu số lượng đồng phân hay khối lượng mol (bộ máy tính toán sẽ thêm riêng). Không bịa số liệu."""


def _wrap(text, width):
    lines, cur = [], ""
    for w in str(text).split():
        if len(cur) + len(w) + (1 if cur else 0) <= width:
            cur = (cur + " " + w).strip()
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines or [""]


def _tint(hexc, a=0.12):
    r, g, b = int(hexc[1:3], 16), int(hexc[3:5], 16), int(hexc[5:7], 16)
    return f"rgba({r},{g},{b},{a})"


def mindmap_layout(title, branches):
    """Tính vị trí mọi khung và đường nối (đơn vị px). Dùng chung cho SVG, PowerPoint, draw.io và FreeMind."""
    LEAF_W, LH = 30, 19
    cw_leaf, cw_br, cw_root = 7.9, 9.8, 14.0
    root_lines = _wrap(title, 16)
    rw = max(len(x) for x in root_lines) * cw_root + 56
    rh = len(root_lines) * 28 + 36
    items = []
    for i, b in enumerate(branches):
        blines = _wrap(b["ten"], 18)
        bw = max(len(x) for x in blines) * cw_br + 36
        bh = len(blines) * 22 + 20
        leaves = []
        for t in b.get("y", []):
            ll = _wrap(t, LEAF_W)
            leaves.append((t, ll, max(len(x) for x in ll) * cw_leaf + 24, len(ll) * LH + 14))
        block = sum(l[3] for l in leaves) + 8 * max(0, len(leaves) - 1)
        items.append({"b": b, "blines": blines, "bw": bw, "bh": bh, "leaves": leaves, "block": block,
                      "h": max(block, bh), "color": b.get("color") or MM_COLORS[i % len(MM_COLORS)]})
    half = (len(items) + 1) // 2
    sides = {"R": items[:half], "L": items[half:]}
    GAP, MARG = 28, 50
    tot = {k: sum(x["h"] for x in v) + GAP * max(0, len(v) - 1) for k, v in sides.items()}
    H = max(max(tot.values()), rh) + 2 * MARG
    reach = 0
    for v in sides.values():
        for x in v:
            reach = max(reach, 70 + x["bw"] + 55 + max([l[2] for l in x["leaves"]] or [0]))
    W = 2 * (rw / 2 + reach + MARG)
    cx, cy = W / 2, H / 2
    nodes = [{"id": "n0", "kind": "root", "x": cx - rw / 2, "y": cy - rh / 2, "w": rw, "h": rh, "text": title,
              "lines": root_lines, "color": "#4f46e5", "side": None, "parent": None}]
    edges, k = [], 1
    for side, v in sides.items():
        sgn = 1 if side == "R" else -1
        y = cy - tot[side] / 2
        for x in v:
            col, yb = x["color"], y + x["h"] / 2
            bx0 = cx + sgn * (rw / 2 + 70)
            bx1 = bx0 + sgn * x["bw"]
            bid = f"n{k}"
            k += 1
            nodes.append({"id": bid, "kind": "branch", "x": min(bx0, bx1), "y": yb - x["bh"] / 2, "w": x["bw"], "h": x["bh"],
                          "text": x["b"]["ten"], "lines": x["blines"], "color": col, "side": side, "parent": "n0"})
            edges.append({"a": "n0", "b": bid, "color": col, "width": 5, "x1": cx + sgn * rw / 2, "y1": cy, "x2": bx0, "y2": yb, "side": side})
            ly = y + (x["h"] - x["block"]) / 2
            lx0 = bx1 + sgn * 55
            for text, ll, lw, lh in x["leaves"]:
                lid = f"n{k}"
                k += 1
                nodes.append({"id": lid, "kind": "leaf", "x": lx0 if side == "R" else lx0 - lw, "y": ly, "w": lw, "h": lh,
                              "text": text, "lines": ll, "color": col, "side": side, "parent": bid})
                edges.append({"a": bid, "b": lid, "color": col, "width": 2.5, "x1": bx1, "y1": yb, "x2": lx0, "y2": ly + lh / 2, "side": side})
                ly += lh + 8
            y += x["h"] + GAP
    return {"W": W, "H": H, "nodes": nodes, "edges": edges}


def svg_from_layout(L):
    esc = html.escape
    W, H = L["W"], L["H"]
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{W:.0f}" height="{H:.0f}" viewBox="0 0 {W:.0f} {H:.0f}" '
           f'font-family="{MM_FONT}"><defs><linearGradient id="rg" x1="0" y1="0" x2="1" y2="1"><stop offset="0" stop-color="#0e8f6e"/>'
           f'<stop offset="1" stop-color="#14b8a6"/></linearGradient></defs><rect width="100%" height="100%" fill="#ffffff"/>']
    for e in L["edges"]:
        xm = (e["x1"] + e["x2"]) / 2
        out.append(f'<path d="M{e["x1"]:.1f} {e["y1"]:.1f} C{xm:.1f} {e["y1"]:.1f} {xm:.1f} {e["y2"]:.1f} {e["x2"]:.1f} {e["y2"]:.1f}" '
                   f'fill="none" stroke="{e["color"]}" stroke-width="{e["width"]}" stroke-linecap="round"/>')
    texts = []
    for n in L["nodes"]:
        x, y, w, h, col = n["x"], n["y"], n["w"], n["h"], n["color"]
        cxn, cyn = x + w / 2, y + h / 2
        if n["kind"] == "root":
            out.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="{h / 2:.1f}" fill="url(#rg)"/>')
            ty = cyn - len(n["lines"]) * 28 / 2 + 21
            for j, ln in enumerate(n["lines"]):
                texts.append(f'<text x="{cxn:.1f}" y="{ty + j * 28:.1f}" text-anchor="middle" font-size="24" font-weight="800" fill="#fff">{esc(ln)}</text>')
        elif n["kind"] == "branch":
            out.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="14" fill="{col}"/>')
            ty = cyn - len(n["lines"]) * 22 / 2 + 16
            for j, ln in enumerate(n["lines"]):
                texts.append(f'<text x="{cxn:.1f}" y="{ty + j * 22:.1f}" text-anchor="middle" font-size="17" font-weight="700" fill="#fff">{esc(ln)}</text>')
        else:
            out.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="10" fill="{_tint(col)}" stroke="{col}" stroke-width="1.5"/>')
            for j, ln in enumerate(n["lines"]):
                texts.append(f'<text x="{x + 12:.1f}" y="{y + 7 + 19 * (j + 1) - 5:.1f}" font-size="15" fill="#1f2937">{esc(ln)}</text>')
    return "".join(out + texts) + "</svg>"


def mindmap_svg(title, branches):
    L = mindmap_layout(title, branches)
    return svg_from_layout(L), L["W"], L["H"]


def _blend(hexc, a=0.12):
    r, g, b = int(hexc[1:3], 16), int(hexc[3:5], 16), int(hexc[5:7], 16)
    return "{:02x}{:02x}{:02x}".format(*(round(c * a + 255 * (1 - a)) for c in (r, g, b)))


def mm_pptx_bytes(L):
    """PowerPoint: mỗi khung là hình có chữ sửa được, đường nối dính vào khung nên kéo khung thì đường nối đi theo."""
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Emu, Pt
    except ImportError:
        return None
    E = 9525
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(L["W"] * E)), Emu(int(L["H"] * E))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rgb = lambda h: RGBColor.from_string(h.lstrip("#").upper())
    shapes = {}
    for n in L["nodes"]:
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(n["x"] * E)), Emu(int(n["y"] * E)), Emu(int(n["w"] * E)), Emu(int(n["h"] * E)))
        shp.adjustments[0] = min(0.5, (n["h"] / 2 if n["kind"] == "root" else (14 if n["kind"] == "branch" else 10)) / min(n["w"], n["h"]))
        shp.shadow.inherit = False
        shp.fill.solid()
        if n["kind"] == "leaf":
            shp.fill.fore_color.rgb = RGBColor.from_string(_blend(n["color"]).upper())
            shp.line.color.rgb = rgb(n["color"])
            shp.line.width = Pt(1.2)
        else:
            shp.fill.fore_color.rgb = rgb(n["color"])
            shp.line.fill.background()
        tf = shp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Emu(int(10 * E))
        tf.margin_top = tf.margin_bottom = Emu(int(4 * E))
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT if n["kind"] == "leaf" else PP_ALIGN.CENTER
        run = para.add_run()
        run.text = n["text"]
        run.font.name = "Arial"
        run.font.size = Pt({"root": 18, "branch": 13, "leaf": 11}[n["kind"]])
        run.font.bold = n["kind"] != "leaf"
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37) if n["kind"] == "leaf" else RGBColor(255, 255, 255)
        shapes[n["id"]] = shp
    tree = slide.shapes._spTree
    for e in L["edges"]:
        c = slide.shapes.add_connector(MSO_CONNECTOR.CURVE, Emu(int(e["x1"] * E)), Emu(int(e["y1"] * E)), Emu(int(e["x2"] * E)), Emu(int(e["y2"] * E)))
        right = e["side"] == "R"
        c.begin_connect(shapes[e["a"]], 3 if right else 1)
        c.end_connect(shapes[e["b"]], 1 if right else 3)
        c.line.color.rgb = rgb(e["color"])
        c.line.width = Pt(e["width"] * 0.75)
        tree.remove(c._element)
        tree.insert(2, c._element)  # đưa đường nối xuống dưới các khung
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def mm_drawio(L):
    """File .drawio (mở bằng app.diagrams.net hoặc draw.io Desktop)."""
    def attr(t):
        return html.escape(html.escape(str(t), quote=False), quote=True)
    cells = ['<mxCell id="0"/>', '<mxCell id="1" parent="0"/>']
    for n in L["nodes"]:
        if n["kind"] == "root":
            st_ = "rounded=1;arcSize=50;whiteSpace=wrap;html=1;fillColor=#4f46e5;strokeColor=none;fontColor=#ffffff;fontStyle=1;fontSize=18;"
        elif n["kind"] == "branch":
            st_ = f"rounded=1;arcSize=30;whiteSpace=wrap;html=1;fillColor={n['color']};strokeColor=none;fontColor=#ffffff;fontStyle=1;fontSize=13;"
        else:
            st_ = (f"rounded=1;arcSize=20;whiteSpace=wrap;html=1;fillColor=#{_blend(n['color'])};strokeColor={n['color']};"
                   "fontColor=#1f2937;fontSize=11;align=left;spacingLeft=6;")
        cells.append(f'<mxCell id="{n["id"]}" value="{attr(n["text"])}" style="{st_}" vertex="1" parent="1">'
                     f'<mxGeometry x="{n["x"]:.0f}" y="{n["y"]:.0f}" width="{n["w"]:.0f}" height="{n["h"]:.0f}" as="geometry"/></mxCell>')
    for i, e in enumerate(L["edges"]):
        ex, nx = ("1", "0") if e["side"] == "R" else ("0", "1")
        cells.append(f'<mxCell id="e{i}" style="curved=1;endArrow=none;html=1;strokeColor={e["color"]};strokeWidth={e["width"]};'
                     f'exitX={ex};exitY=0.5;entryX={nx};entryY=0.5;" edge="1" parent="1" source="{e["a"]}" target="{e["b"]}">'
                     '<mxGeometry relative="1" as="geometry"/></mxCell>')
    return ('<mxfile host="app.diagrams.net"><diagram name="So do tu duy" id="mm1"><mxGraphModel grid="1" gridSize="10" page="0">'
            "<root>" + "".join(cells) + "</root></mxGraphModel></diagram></mxfile>")


def mm_freemind(L):
    """File .mm (FreeMind) mở được bằng XMind, Freeplane, MindManager, Mindomo."""
    q = lambda t: html.escape(str(t), quote=True)
    root = L["nodes"][0]
    out = [f'<map version="1.0.1"><node TEXT="{q(root["text"])}">']
    for b in [n for n in L["nodes"] if n["kind"] == "branch"]:
        pos = "right" if b["side"] == "R" else "left"
        out.append(f'<node TEXT="{q(b["text"])}" POSITION="{pos}" COLOR="{b["color"]}">')
        out += [f'<node TEXT="{q(l["text"])}"/>' for l in L["nodes"] if l["parent"] == b["id"]]
        out.append("</node>")
    out.append("</node></map>")
    return "".join(out)


def rows_to_branches(rows):
    """Bảng chỉnh sửa (Nhánh, Ý) -> danh sách nhánh. Dòng không ghi nhánh thì thuộc nhánh ngay trên."""
    order, data = [], {}
    for r_ in rows:
        n = str(r_.get("Nhánh") or "").strip()
        y = str(r_.get("Ý") or "").strip()
        if not n and not y:
            continue
        n = n or (order[-1] if order else "Nhánh mới")
        if n not in data:
            data[n] = []
            order.append(n)
        if y:
            data[n].append(y[:100])
    return [{"ten": n[:28], "y": data[n], **({"color": "#475569"} if n.startswith("Số liệu (máy tính)") else {})} for n in order[:10]]


def mm_outline(title, branches):
    lines = [f"# {title}"]
    for b in branches:
        lines.append(f"\n## {b['ten']}")
        lines += [f"- {t}" for t in b.get("y", [])]
    return "\n".join(lines)


def mm_clean(raw):
    """Giữ sơ đồ gọn: tối đa 8 nhánh, 5 ý mỗi nhánh, cắt câu quá dài."""
    out = []
    for b in (raw.get("nhanh") or [])[:8]:
        ten = str(b.get("ten", "")).strip()[:28]
        ys = [str(t).strip()[:100] for t in (b.get("y") or []) if str(t).strip()][:5]
        if ten and ys:
            out.append({"ten": ten, "y": ys})
    return out


def mm_engine_facts(smiles):
    """Số liệu do RDKit và bộ máy đồng phân tính, không phải AI."""
    m, _ = parse_student(smiles or "")
    if m is None:
        return None
    f = formula_of(m)
    facts = [f"Công thức phân tử: {pretty(f)}", f"Khối lượng mol: {Descriptors.MolWt(m):.2f} g/mol",
             "Nhóm chức / loại: " + ", ".join(iso_labels(m))]
    counts, nH = mol_counts(m)
    if sum(counts.values()) <= 10 and all(e in VALENCE for e in counts):
        try:
            res = build_iso_result(hill_key(counts, nH), None, dict(PRESETS["mono"]))
            if res["ok"] and res["kept"]:
                k = res["kept"]
                facts.append(f"Số đồng phân cấu tạo đơn chức mạch hở của {pretty(f)}: {len(k)} (kể cả lập thể: {sum(x['stereo'] for x in k)})")
        except Exception:
            pass
    return facts


def mm_component(svg, W, H):
    scale = min(1.0, 1000.0 / W)
    page = ('<div style="font-family:Arial,sans-serif"><button id="dl" style="background:#1f7a4d;color:#fff;border:0;border-radius:10px;'
            'padding:9px 16px;font-weight:700;cursor:pointer">⬇️ Tải ảnh PNG</button>'
            '<div style="overflow:auto;border:1px solid #d6ece5;border-radius:12px;margin-top:8px;background:#fff">'
            '<div style="width:max-content;max-width:100%">' + svg.replace("<svg ", '<svg style="max-width:100%;height:auto;display:block" ', 1) +
            '</div></div><script>document.getElementById("dl").onclick=function(){'
            'var s=document.querySelector("svg"),x=new XMLSerializer().serializeToString(s),i=new Image();'
            'i.onload=function(){var c=document.createElement("canvas");c.width=' + str(int(W * 2)) + ';c.height=' + str(int(H * 2)) + ';'
            'var g=c.getContext("2d");g.fillStyle="#fff";g.fillRect(0,0,c.width,c.height);g.drawImage(i,0,0,c.width,c.height);'
            'c.toBlob(function(b){var a=document.createElement("a");a.href=URL.createObjectURL(b);a.download="so_do_tu_duy.png";a.click();});};'
            'i.src="data:image/svg+xml;charset=utf-8,"+encodeURIComponent(x);};</script></div>')
    height = int(H * scale) + 90
    if hasattr(st, "iframe"):  # st.components.v1.html sẽ bị gỡ; nội dung trong page đã được escape, không chứa HTML từ AI
        st.iframe(page, height=height)
    else:
        components.html(page, height=height, scrolling=True)


MM_EXAMPLES = ["Ancol", "Anđehit và xeton", "Axit cacboxylic", "Este và lipit", "Amin và amino axit", "Benzen và đồng đẳng",
               "Phenol", "Etanol", "Phản ứng cộng"]


def _set_topic(t):
    st.session_state["mm_topic"] = t


def mindmap_ui():
    top = st.container()
    st.session_state.setdefault("mm", None)
    st.text_input("Chủ đề hoặc tên chất", key="mm_topic", placeholder="Ví dụ: Ancol, Etanol, Phản ứng thế, Este và lipit")
    cols = st.columns(len(MM_EXAMPLES))
    for c, t in zip(cols, MM_EXAMPLES):
        c.button(t, key=f"mm_ex_{t}", on_click=_set_topic, args=(t,), width="stretch")
    level = st.radio("Mức chi tiết", ["Gọn (4 nhánh, mỗi nhánh 2-3 ý)", "Đầy đủ (6 nhánh, mỗi nhánh 3-4 ý)"], horizontal=True, key="mm_level")
    topic = st.session_state.get("mm_topic", "").strip()
    if st.button("Tạo sơ đồ tư duy", type="primary", disabled=not topic, key="mm_go"):
        if st.session_state.n >= MAX_PER_SESSION:
            show_top_error(top, "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục.")
        else:
            with st.spinner("Đang vẽ sơ đồ…"):
                try:
                    want = "4 nhánh, mỗi nhánh 2-3 ý" if level.startswith("Gọn") else "6 nhánh, mỗi nhánh 3-4 ý"
                    raw = ask_ai([MM_PROMPT + f"\nSố nhánh và ý: {want}.\n\nChủ đề: {topic}"], temperature=0.2)
                    branches = mm_clean(raw)
                    if not branches:
                        show_top_error(top, "Chưa tạo được sơ đồ cho chủ đề này, hãy thử diễn đạt rõ hơn.")
                    else:
                        smi, facts = (raw.get("chat_smiles") or "").strip(), None
                        if smi and smi.lower() != "null":
                            facts = mm_engine_facts(smi)
                        if facts:
                            branches.append({"ten": "Số liệu (máy tính)", "y": facts, "color": "#475569"})
                        st.session_state.mm = {"title": str(raw.get("chu_de") or topic)[:60], "branches": branches,
                                               "smiles": smi if facts else "", "rev": int(time.time() * 1000)}
                        add_history(st.session_state.mm['title'], 'Sơ đồ')
                        st.session_state.n += 1
                except Exception as e:
                    st.session_state.mm = None
                    show_top_error(top, "Có lỗi khi tạo sơ đồ, hãy thử lại.", e)
    r = st.session_state.get("mm")
    if r:
        st.divider()
        rev = r.setdefault("rev", 0)
        title = st.text_input("Chủ đề ở giữa sơ đồ", value=r["title"], key=f"mm_title_{rev}").strip() or r["title"]
        with st.expander("✏️ Chỉnh sửa nội dung (sửa chữ, thêm hoặc xóa dòng, sơ đồ cập nhật ngay)", expanded=False):
            base = [{"Nhánh": b["ten"], "Ý": t} for b in r["branches"] for t in (b["y"] or [""])]
            edited = st.data_editor(base, num_rows="dynamic", hide_index=True, key=f"mm_edit_{rev}", width="stretch",
                                    column_config={"Nhánh": st.column_config.TextColumn("Nhánh", width="medium"),
                                                   "Ý": st.column_config.TextColumn("Ý", width="large")})
            st.caption("Muốn thêm nhánh mới: thêm dòng mới và gõ tên nhánh mới. Dòng để trống cột Nhánh sẽ thuộc nhánh ngay phía trên.")
        rows = edited.to_dict("records") if hasattr(edited, "to_dict") else list(edited)
        branches = rows_to_branches(rows)
        if not branches:
            st.warning("Sơ đồ đang trống, hãy nhập ít nhất một nhánh.")
        else:
            L = mindmap_layout(title, branches)
            svg = svg_from_layout(L)
            mm_component(svg, L["W"], L["H"])
            st.markdown("**Tải về để chỉnh sửa tiếp trên máy**")
            c1, c2, c3 = st.columns(3)
            pp = mm_pptx_bytes(L)
            if pp:
                c1.download_button("📊 PowerPoint (.pptx)", pp, file_name="so_do_tu_duy.pptx", key="mm_pptx",
                                   mime="application/vnd.openxmlformats-officedocument.presentationml.presentation", width="stretch")
            else:
                c1.caption("Thiếu thư viện python-pptx nên chưa xuất được PowerPoint.")
            c2.download_button("🧩 draw.io (.drawio)", mm_drawio(L).encode("utf-8"), file_name="so_do_tu_duy.drawio", key="mm_dio", width="stretch")
            c3.download_button("🗺️ XMind / FreeMind (.mm)", mm_freemind(L).encode("utf-8"), file_name="so_do_tu_duy.mm", key="mm_fm", width="stretch")
            d1, d2 = st.columns(2)
            d1.download_button("🖼️ SVG (vector, in sắc nét)", svg.encode("utf-8"), file_name="so_do_tu_duy.svg", mime="image/svg+xml", key="mm_svg", width="stretch")
            d2.download_button("📝 Dàn ý (Markdown)", mm_outline(title, branches).encode("utf-8"), file_name="so_do_tu_duy.md", key="mm_md", width="stretch")
            st.caption("PowerPoint: mỗi khung là hình có chữ sửa được và đường nối dính vào khung, kéo khung thì đường nối đi theo. "
                       "draw.io mở bằng app.diagrams.net. File .mm mở bằng XMind, Freeplane, MindManager.")
        if r.get("smiles"):
            with st.expander("Cấu trúc chất (kiểm tra xem đúng chất bạn muốn chưa)"):
                m = Chem.MolFromSmiles(r["smiles"])
                if m is not None:
                    st.image(Draw.MolToImage(m, size=(300, 200)))
                st.code(r["smiles"], language=None)
        st.caption("Nội dung lý thuyết do AI soạn theo SGK nên cần đối chiếu sách. Nhánh \"Số liệu (máy tính)\" do RDKit và bộ máy đếm đồng phân tính, không phải AI.")


# ---------------- Giao diện ----------------
# ====================================================================
#  GIAO DIỆN MolSnap (thay thế lớp giao diện cũ, giữ nguyên toàn bộ logic phía trên)
# ====================================================================
def _uri(svg):
    return "data:image/svg+xml;base64," + base64.b64encode(svg.encode("utf-8")).decode()


# ---------- Bộ biểu tượng nét (24x24) ----------
ICON_PATHS = {
    "home": "<path d='M3 11l9-8 9 8'/><path d='M5 10v10h5v-6h4v6h5V10'/>",
    "structure": "<circle cx='6' cy='6' r='2'/><circle cx='18' cy='6' r='2'/><circle cx='12' cy='12' r='2'/><circle cx='6' cy='18' r='2'/><circle cx='18' cy='18' r='2'/><path d='M7.5 7.5l3 3M16.5 7.5l-3 3M7.5 16.5l3-3M16.5 16.5l-3-3'/>",
    "flask": "<path d='M9 3h6M10 3v6l-5 9a2 2 0 0 0 2 3h10a2 2 0 0 0 2-3l-5-9V3'/><path d='M7.5 15h9'/>",
    "isomer": "<path d='M16 3h5v5M4 20L21 3M21 16v5h-5M15 15l6 6M4 4l5 5'/>",
    "equipment": "<path d='M7 3h10v3H7zM6 6h12v14a1 1 0 0 1-1 1H7a1 1 0 0 1-1-1z'/><path d='M9 12h6M9 16h6'/>",
    "mind": "<circle cx='18' cy='5' r='3'/><circle cx='6' cy='12' r='3'/><circle cx='18' cy='19' r='3'/><path d='M8.6 13.5l6.8 4M15.4 6.5l-6.8 4'/>",
    "practice": "<rect x='5' y='4' width='14' height='17' rx='2'/><path d='M9 4h6v3H9zM9 14l2 2 4-4'/>",
    "stats": "<path d='M5 21V11M12 21V4M19 21v-7'/>",
    "chat": "<path d='M21 12a8 8 0 0 1-11.5 7.2L4 21l1.8-5.3A8 8 0 1 1 21 12z'/>",
    "pen": "<path d='M12 20h9'/><path d='M16.5 3.5a2.1 2.1 0 0 1 3 3L7 19l-4 1 1-4z'/>",
    "search": "<circle cx='11' cy='11' r='7'/><path d='M21 21l-4.3-4.3'/>",
    "sun": "<circle cx='12' cy='12' r='4'/><path d='M12 2v2M12 20v2M4.9 4.9l1.4 1.4M17.7 17.7l1.4 1.4M2 12h2M20 12h2M4.9 19.1l1.4-1.4M17.7 6.3l1.4-1.4'/>",
    "bell": "<path d='M6 8a6 6 0 0 1 12 0c0 7 3 9 3 9H3s3-2 3-9'/><path d='M10.3 21a1.9 1.9 0 0 0 3.4 0'/>",
    "image": "<rect x='3' y='3' width='18' height='18' rx='3'/><circle cx='9' cy='9' r='2'/><path d='M21 15l-5-5L5 21'/>",
    "arrow": "<path d='M5 12h14M13 6l6 6-6 6'/>",
    "camera": "<path d='M4 8h3l2-3h6l2 3h3v11H4z'/><circle cx='12' cy='13' r='3.5'/>",
}


def icon_svg(name, color="currentColor", size=22, sw=1.8):
    return (f"<svg xmlns='http://www.w3.org/2000/svg' width='{size}' height='{size}' viewBox='0 0 24 24' fill='none' stroke='{color}' "
            f"stroke-width='{sw}' stroke-linecap='round' stroke-linejoin='round'>{ICON_PATHS[name]}</svg>")


def _mask(name):
    return f'url("{_uri(icon_svg(name, "black"))}")'


# ---------- Hình minh họa ----------
LOGO_SVG = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 64 64"><defs>
<linearGradient id="lg" x1="0" y1="0" x2="1" y2="1"><stop offset="0" stop-color="#8b5cf6"/><stop offset=".55" stop-color="#4f6df5"/><stop offset="1" stop-color="#22d3ee"/></linearGradient></defs>
<path d="M32 5 L56 19 V45 L32 59 L8 45 V19 Z" fill="url(#lg)" stroke="url(#lg)" stroke-width="6" stroke-linejoin="round"/>
<path d="M42 23 C37 16 23 18 23 26 C23 34 41 31 41 40 C41 48 28 48 22 41" fill="none" stroke="#fff" stroke-width="5.5" stroke-linecap="round"/>
<circle cx="45" cy="21" r="3" fill="#fff" opacity=".9"/><circle cx="19" cy="43" r="3" fill="#fff" opacity=".9"/></svg>"""

HERO_ART = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 560 290"><defs>
<linearGradient id="liq" x1="0" y1="0" x2="0" y2="1"><stop offset="0" stop-color="#67e8f9"/><stop offset=".55" stop-color="#3b82f6"/><stop offset="1" stop-color="#4f46e5"/></linearGradient>
<linearGradient id="glass" x1="0" y1="0" x2="1" y2="1"><stop offset="0" stop-color="#fff" stop-opacity=".55"/><stop offset="1" stop-color="#bfdbfe" stop-opacity=".25"/></linearGradient>
<linearGradient id="leaf" x1="0" y1="0" x2="1" y2="1"><stop offset="0" stop-color="#86efac"/><stop offset="1" stop-color="#16a34a"/></linearGradient>
<style>.f1{animation:f 5s ease-in-out infinite}.f2{animation:f 6.5s ease-in-out 1s infinite}.f3{animation:f 7.5s ease-in-out .4s infinite}
@keyframes f{0%,100%{transform:translateY(0)}50%{transform:translateY(-8px)}}</style></defs>
<path d="M395 250 Q360 190 400 150 Q430 190 395 250Z" fill="url(#leaf)" opacity=".95"/>
<path d="M410 252 Q420 190 470 170 Q470 220 410 252Z" fill="#22c55e" opacity=".9"/>
<path d="M340 214 Q388 192 442 206 L442 266 Q388 252 340 272Z" fill="#ece9ff"/>
<path d="M442 206 Q500 190 552 212 L552 270 Q500 252 442 266Z" fill="#fff"/>
<path d="M442 206 V266" stroke="#c7c9f5" stroke-width="2"/>
<g stroke="#c4c8ef" stroke-width="2.5" stroke-linecap="round"><path d="M462 232 H520"/><path d="M462 244 H532"/><path d="M356 232 H420"/><path d="M356 244 H410"/></g>
<text x="466" y="226" font-family="Dancing Script,cursive" font-size="20" fill="#7a86c9" font-style="italic">H₂O</text>
<polygon points="378,226 388,220 398,226 398,238 388,244 378,238" fill="none" stroke="#7a86c9" stroke-width="2"/>
<g class="f1"><path d="M290 34 H322 V98 H290Z" fill="url(#glass)" stroke="#fff" stroke-opacity=".8" stroke-width="2"/>
<ellipse cx="306" cy="32" rx="24" ry="8" fill="none" stroke="#a5f3fc" stroke-width="6"/>
<path d="M290 98 H322 L372 196 Q384 230 352 232 H260 Q228 230 240 196Z" fill="url(#glass)" stroke="#fff" stroke-opacity=".85" stroke-width="2"/>
<path d="M254 170 H358 L372 196 Q384 230 352 232 H260 Q228 230 240 196Z" fill="url(#liq)"/>
<path d="M254 170 Q306 160 358 170" fill="none" stroke="#e0f2fe" stroke-width="3" opacity=".8"/>
<circle cx="276" cy="198" r="7" fill="#e0f2fe" opacity=".75"/><circle cx="312" cy="190" r="5" fill="#e0f2fe" opacity=".75"/><circle cx="332" cy="210" r="6" fill="#e0f2fe" opacity=".75"/><circle cx="300" cy="218" r="4" fill="#e0f2fe" opacity=".7"/>
<path d="M262 120 Q250 170 252 196" fill="none" stroke="#fff" stroke-width="5" stroke-linecap="round" opacity=".55" transform="translate(30 0)"/></g>
<g class="f2"><g stroke="#cfd8ff" stroke-width="5" stroke-linecap="round"><line x1="70" y1="132" x2="128" y2="88"/><line x1="128" y1="88" x2="190" y2="134"/><line x1="190" y1="134" x2="152" y2="186"/><line x1="128" y1="88" x2="214" y2="66"/><line x1="70" y1="132" x2="98" y2="186"/></g>
<circle cx="70" cy="132" r="15" fill="#4f6df5"/><circle cx="128" cy="88" r="19" fill="#8b5cf6"/><circle cx="190" cy="134" r="15" fill="#22d3ee"/><circle cx="152" cy="186" r="17" fill="#3b82f6"/><circle cx="214" cy="66" r="11" fill="#f59e0b"/><circle cx="98" cy="186" r="9" fill="#f97316"/>
<g fill="#fff" opacity=".6"><circle cx="65" cy="126" r="5"/><circle cx="122" cy="80" r="6"/><circle cx="185" cy="129" r="5"/><circle cx="146" cy="179" r="5"/><circle cx="211" cy="63" r="3.5"/></g></g>
<g class="f3"><ellipse cx="232" cy="38" rx="30" ry="12" fill="none" stroke="#7dd3fc" stroke-width="7" opacity=".85" transform="rotate(-14 232 38)"/>
<ellipse cx="452" cy="96" rx="20" ry="8" fill="none" stroke="#c4b5fd" stroke-width="6" opacity=".8" transform="rotate(12 452 96)"/></g>
<g fill="#fff" opacity=".9"><path d="M40 52 l3 9 9 3 -9 3 -3 9 -3-9 -9-3 9-3z"/><path d="M500 40 l2 6 6 2 -6 2 -2 6 -2-6 -6-2 6-2z"/><path d="M360 262 l2 5 5 2 -5 2 -2 5 -2-5 -5-2 5-2z" opacity=".6"/></g></svg>"""

MASCOT_SVG = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 300 250"><defs>
<linearGradient id="lf" x1="0" y1="0" x2="1" y2="1"><stop offset="0" stop-color="#4ade80"/><stop offset="1" stop-color="#15803d"/></linearGradient>
<linearGradient id="rb" x1="0" y1="0" x2="0" y2="1"><stop offset="0" stop-color="#fff"/><stop offset="1" stop-color="#dbe4ff"/></linearGradient>
<style>.bob{animation:b 4s ease-in-out infinite;transform-origin:150px 200px}@keyframes b{0%,100%{transform:translateY(0)}50%{transform:translateY(-5px)}}</style></defs>
<path d="M8 232 Q-6 180 40 160 Q52 205 8 232Z" fill="url(#lf)" opacity=".9"/><path d="M40 238 Q40 190 84 176 Q90 220 40 238Z" fill="#22c55e" opacity=".85"/>
<path d="M290 236 Q306 196 262 176 Q250 214 290 236Z" fill="url(#lf)" opacity=".9"/>
<rect x="50" y="206" width="200" height="26" rx="6" fill="#3b5bdb"/><rect x="50" y="206" width="200" height="8" rx="4" fill="#6c8cff"/>
<rect x="64" y="182" width="176" height="26" rx="6" fill="#7c3aed"/><rect x="64" y="182" width="176" height="8" rx="4" fill="#a78bfa"/>
<rect x="80" y="160" width="150" height="24" rx="6" fill="#0891b2"/><rect x="80" y="160" width="150" height="7" rx="3.5" fill="#67e8f9"/>
<g class="bob"><line x1="150" y1="30" x2="150" y2="48" stroke="#a5b4fc" stroke-width="4"/><circle cx="150" cy="26" r="7" fill="#22d3ee"/>
<rect x="104" y="46" width="92" height="70" rx="30" fill="url(#rb)" stroke="#c7d2fe" stroke-width="3"/>
<rect x="114" y="60" width="72" height="42" rx="20" fill="#16204f"/><ellipse cx="136" cy="80" rx="7" ry="9" fill="#38e1ff"/><ellipse cx="166" cy="80" rx="7" ry="9" fill="#38e1ff"/>
<path d="M143 92 Q151 98 159 92" fill="none" stroke="#38e1ff" stroke-width="3" stroke-linecap="round"/>
<rect x="98" y="66" width="10" height="28" rx="5" fill="#7c8cff"/><rect x="192" y="66" width="10" height="28" rx="5" fill="#7c8cff"/>
<rect x="116" y="116" width="68" height="46" rx="20" fill="url(#rb)" stroke="#c7d2fe" stroke-width="3"/>
<circle cx="150" cy="138" r="6" fill="#6d5bf0"/>
<path d="M118 126 Q96 140 104 158" fill="none" stroke="#dbe4ff" stroke-width="12" stroke-linecap="round"/>
<path d="M182 126 Q204 132 208 148" fill="none" stroke="#dbe4ff" stroke-width="12" stroke-linecap="round"/>
<g transform="translate(196 134)"><path d="M10 0 h10 v14 l12 22 q4 10 -6 10 h-22 q-10 0 -6 -10 l12 -22z" fill="#dff7ff" stroke="#fff" stroke-width="2"/><path d="M4 30 h32 l4 8 q3 8 -6 8 h-28 q-9 0 -6 -8z" fill="#4ade80"/></g></g></svg>"""

AI_FACE = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 64 64"><rect x="8" y="12" width="48" height="40" rx="18" fill="#fff"/><rect x="14" y="20" width="36" height="24" rx="12" fill="#16204f"/><ellipse cx="25" cy="32" rx="4" ry="5" fill="#38e1ff"/><ellipse cx="39" cy="32" rx="4" ry="5" fill="#38e1ff"/><circle cx="32" cy="8" r="4" fill="#22d3ee"/><rect x="31" y="9" width="2" height="5" fill="#a5b4fc"/></svg>"""

PAPER_SVG = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 150 130"><defs><filter id="s" x="-10%" y="-10%" width="130%" height="130%"><feDropShadow dx="0" dy="3" stdDeviation="3" flood-color="#1b2150" flood-opacity=".25"/></filter></defs>
<g transform="rotate(4 75 65)"><rect x="18" y="14" width="112" height="100" rx="3" fill="#dcd5c8" filter="url(#s)"/></g>
<g transform="rotate(-3 75 65)"><rect x="14" y="10" width="112" height="100" rx="3" fill="#f6f3ec" filter="url(#s)"/>
<g fill="none" stroke="#5b6075" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round">
<polygon points="52,58 64,51 76,58 76,72 64,79 52,72"/><polyline points="55,60 55,70" stroke-width="1.6"/><polyline points="66,77 74,72" stroke-width="1.6" transform="translate(0 0)"/><polyline points="58,54 64,51 70,54" stroke-width="1.6"/>
<line x1="76" y1="58" x2="88" y2="51"/><line x1="88" y1="51" x2="88" y2="38"/><line x1="91" y1="51" x2="91" y2="38"/><line x1="88" y1="51" x2="100" y2="58"/></g>
<text x="86" y="33" font-family="Dancing Script,cursive" font-size="13" fill="#5b6075">O</text><text x="100" y="66" font-family="Dancing Script,cursive" font-size="13" fill="#5b6075">OH</text></g></svg>"""


# ---------- Tài nguyên giao diện ----------
CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Be+Vietnam+Pro:wght@400;500;600;700;800&family=Dancing+Script:wght@600;700&display=swap');
:root{--pri:#4f5cf0;--vio:#7c4dff;--cy:#19c3ff;--ink:#1b2150;--mut:#6b7299;--line:#e3e7fb;--soft:#f3f5ff;--navy:#0a0e30;
--grad:linear-gradient(135deg,#3b6cf6 0%,#6d4df0 100%)}
html{font-size:__ROOT__!important}
html,body,.stApp,.stMarkdown,p,label,input,textarea,button,li,h1,h2,h3,h4{font-family:'Be Vietnam Pro',system-ui,sans-serif}
.stApp{background:radial-gradient(64.286rem 30rem at 78% -8%,#e3e8ff 0%,rgba(227,232,255,0) 70%),#f5f7ff}
[data-testid="stHeader"]{background:transparent}
[data-testid="stAppDeployButton"]{display:none}
.block-container{max-width:95.714rem;padding:.6rem 1.6rem .8rem}
.block-container [data-testid="stVerticalBlock"]{gap:.7rem}
[data-testid="stSidebarUserContent"] [data-testid="stVerticalBlock"]{gap:.35rem}

/* ---------- Sidebar ---------- */
[data-testid="stSidebar"]{background:linear-gradient(185deg,#080c2c 0%,#0e1450 52%,#0a0f3a 100%);border-right:0}
[data-testid="stSidebarContent"]{background:transparent}
[data-testid="stSidebarUserContent"]{padding:0 .9rem .5rem}
[data-testid="stSidebar"] *{color:#e6eaff}
[data-testid="stSidebar"] [data-testid="stSidebarCollapseButton"] *{color:#c3cbff}
.sb-brand{display:flex;align-items:center;gap:0.857rem;margin:0 0 .6rem}
.sb-brand img{width:3rem;height:3rem;flex:0 0 auto;filter:drop-shadow(0 0.429rem 1rem rgba(110,90,255,.55))}
.sb-brand b{display:block;font-size:1.7rem;font-weight:800;letter-spacing:-.5px;line-height:1.05;color:#fff}
.sb-brand b i{font-style:normal;background:linear-gradient(90deg,#6cb8ff,#a78bfa);-webkit-background-clip:text;background-clip:text;color:transparent}
.sb-brand span{display:block;white-space:nowrap;font-style:italic;font-size:.72rem;color:#aab4ea;margin-top:0.214rem}
.st-key-sidenav [role="radiogroup"]{display:flex;flex-direction:column;gap:0.214rem;width:100%}
.st-key-sidenav [role="radiogroup"]>div{width:100%}
.st-key-sidenav [data-testid="stRadioOption"]>div{flex:1;min-width:0}
.st-key-sidenav [data-testid="stRadioOption"] p{white-space:nowrap}
.st-key-sidenav :is(label[data-baseweb="radio"],[data-testid="stRadioOption"]){display:flex;align-items:center;gap:0.857rem;margin:0;width:100%;padding:0.5rem 1rem;border-radius:0.857rem;cursor:pointer;border:1px solid transparent;transition:background .15s,transform .15s}
.st-key-sidenav :is(label[data-baseweb="radio"],[data-testid="stRadioOption"])>div:first-child{display:none}
.st-key-sidenav [data-testid="stRadioOption"]>div{display:flex;align-items:center;gap:0;width:auto;padding:0}
.st-key-sidenav [data-testid="stRadioOption"]>div>div:first-child{display:none}
.st-key-sidenav :is(label[data-baseweb="radio"],[data-testid="stRadioOption"])::before{content:"";flex:0 0 1.571rem;width:1.571rem;height:1.571rem;background-color:#dbe1ff;-webkit-mask:var(--ic) center/contain no-repeat;mask:var(--ic) center/contain no-repeat}
.st-key-sidenav :is(label[data-baseweb="radio"],[data-testid="stRadioOption"]) p{font-size:1rem;font-weight:500;color:#e6eaff}
.st-key-sidenav :is(label[data-baseweb="radio"],[data-testid="stRadioOption"]):hover{background:rgba(255,255,255,.08);transform:translateX(2px)}
.st-key-sidenav :is(label[data-baseweb="radio"]:has(input:checked),[data-testid="stRadioOption"][data-selected="true"]){background:linear-gradient(90deg,#3f63ee,#5b86ff);border-color:rgba(255,255,255,.28);box-shadow:0 0.571rem 1.571rem rgba(63,99,238,.45)}
.st-key-sidenav :is(label[data-baseweb="radio"]:has(input:checked),[data-testid="stRadioOption"][data-selected="true"])::before{background-color:#fff}
.st-key-sidenav :is(label[data-baseweb="radio"]:has(input:checked),[data-testid="stRadioOption"][data-selected="true"]) p{color:#fff;font-weight:700}
__NAVICONS__
.aicard{display:flex;align-items:center;gap:0.857rem;padding:0.643rem 0.857rem;border-radius:1.143rem;background:linear-gradient(135deg,rgba(63,99,238,.42),rgba(109,77,240,.32));border:1px solid rgba(130,160,255,.5);box-shadow:0 0.714rem 1.857rem rgba(10,14,60,.45);margin-top:0.429rem}
.aicard img{width:2.714rem;height:2.714rem;flex:0 0 auto}
.aicard b{display:block;color:#fff;font-size:1rem}.aicard span{display:block;font-size:.78rem;color:#c9d2ff;line-height:1.35;margin-top:2px}
.aicard .go{margin-left:auto;flex:0 0 auto;width:2.143rem;height:2.143rem;border-radius:50%;background:rgba(255,255,255,.16);display:grid;place-items:center;color:#fff;font-weight:700}
.mascot{position:relative;margin:0 auto;width:72%}.mascot img{width:100%;display:block}
.mascot .hw{position:absolute;right:0.429rem;top:1rem;font-family:'Dancing Script',cursive;font-size:1.35rem;line-height:1.1;transform:rotate(-8deg);color:#e8ecff;text-align:center}

/* ---------- Ô bấm phủ lên thẻ (thẻ HTML + nút ẩn) ---------- */
div[class*="st-key-qk_"],.st-key-aibtn{position:relative;gap:0}
div[class*="st-key-qk_"] [data-testid="stElementContainer"]:has([data-testid="stButton"]),.st-key-aibtn [data-testid="stElementContainer"]:has([data-testid="stButton"]){position:absolute;inset:0;z-index:3;width:100%!important}
div[class*="st-key-qk_"] [data-testid="stButton"],.st-key-aibtn [data-testid="stButton"],div[class*="st-key-qk_"] [data-testid="stButton"] button,.st-key-aibtn [data-testid="stButton"] button{width:100%;height:100%;min-height:0}
div[class*="st-key-qk_"] button,.st-key-aibtn button{opacity:0;cursor:pointer}
div[class*="st-key-qk_"]:hover .tile{transform:translateY(-3px);box-shadow:0 0.857rem 1.857rem rgba(79,92,240,.2)}
.st-key-aibtn:hover .aicard{border-color:#9db4ff;transform:translateY(-1px)}.aicard{transition:all .15s}

/* ---------- Thanh trên ---------- */
.st-key-gsearch [data-testid="InputInstructions"]{display:none}
.st-key-gsearch div[data-baseweb="input"]{background:#fff;border:1px solid var(--line);border-radius:71.357rem;box-shadow:0 0.429rem 1.429rem rgba(79,92,240,.08);height:2.5rem}
.st-key-gsearch input{padding-left:3rem;background:url("__SEARCHIC__") no-repeat 1.1rem center/1.429rem;color:var(--ink)}
.tb-right{display:flex;justify-content:flex-end;align-items:center;gap:1.143rem;color:#5b6293}
.tb-right .av{width:2.571rem;height:2.571rem;border-radius:50%;background:var(--grad);color:#fff;display:grid;place-items:center;font-weight:700;box-shadow:0 0.286rem 0.857rem rgba(79,92,240,.4)}

/* ---------- Hero ---------- */
.hero3{position:relative;display:flex;justify-content:space-between;align-items:center;gap:0.714rem;color:#fff;border-radius:1.714rem;padding:1.143rem 2rem;margin:0 0 0.714rem;overflow:hidden;
background:radial-gradient(30rem 18.571rem at 88% 18%,rgba(25,195,255,.38),rgba(25,195,255,0) 70%),radial-gradient(27.143rem 18.571rem at 40% 110%,rgba(139,92,246,.5),rgba(139,92,246,0) 70%),linear-gradient(115deg,#152a9a 0%,#2f35c8 48%,#5a3bea 100%);
box-shadow:0 1.286rem 3.143rem rgba(47,53,200,.32)}
.hero3 .pill{display:inline-block;font-size:.85rem;background:rgba(255,255,255,.16);border:1px solid rgba(255,255,255,.34);border-radius:71.357rem;padding:0.357rem 1.143rem;font-weight:600;font-size:.95rem;margin-bottom:0.571rem}
.hero3 h1{margin:0;padding:0;font-size:2.05rem;line-height:1.16;font-weight:800;color:#fff;letter-spacing:-.4px}
.hero3 h1 em{font-style:normal;font-size:2.6rem;background:linear-gradient(90deg,#22d3ee,#60a5fa 60%,#a5b4fc);-webkit-background-clip:text;background-clip:text;color:transparent}
.hero3 .sub{margin:0.429rem 0 0.714rem;font-size:1rem;font-weight:600;display:flex;gap:1rem;align-items:center;flex-wrap:wrap}.hero3 .sub i{opacity:.7;font-style:normal}
.hero3 .chips{display:flex;flex-wrap:wrap;gap:0.714rem}
.hero3 .chips span{background:rgba(255,255,255,.14);border:1px solid rgba(255,255,255,.3);border-radius:71.357rem;padding:0.429rem 0.929rem;font-size:.85rem;font-weight:600;backdrop-filter:blur(0.429rem)}
.hero3 .hr{position:relative;flex:0 0 31%}.hero3 .art{width:100%;display:block}
.hero3 .slogan{position:absolute;right:0.714rem;top:0;font-family:'Dancing Script',cursive;font-weight:700;font-size:1.45rem;line-height:1.15;transform:rotate(-7deg);text-align:right;opacity:.96}
.hero3 .slogan u{display:block;height:2px;margin-top:0.429rem;background:linear-gradient(90deg,rgba(255,255,255,0),#fff);border-radius:2px}

/* ---------- Banner trang con ---------- */
.pbanner{display:flex;align-items:center;gap:1rem;color:#fff;border-radius:1.429rem;padding:1rem 1.429rem;margin:0.429rem 0 1.143rem;
background:radial-gradient(21.429rem 11.429rem at 92% 0%,rgba(25,195,255,.35),rgba(25,195,255,0) 70%),linear-gradient(115deg,#152a9a 0%,#3a37cf 55%,#5a3bea 100%);box-shadow:0 0.857rem 2.143rem rgba(47,53,200,.25);animation:fu .45s ease}
.pbanner .pb-ic{width:3.857rem;height:3.857rem;border-radius:1.143rem;background:rgba(255,255,255,.16);border:1px solid rgba(255,255,255,.3);display:grid;place-items:center;flex:0 0 auto}
.pbanner b{display:block;font-size:1.35rem;font-weight:800;color:#fff}.pbanner span{display:block;color:#dbe3ff;font-size:.95rem;margin-top:2px}
@keyframes fu{from{opacity:0;transform:translateY(0.571rem)}to{opacity:1;transform:none}}
.st-key-panel{background:#fff;border:1px solid var(--line);border-radius:1.714rem;padding:1.571rem 1.857rem;box-shadow:0 0.714rem 2.143rem rgba(79,92,240,.08)}
.st-key-panel [role="radiogroup"]{gap:0.571rem;flex-wrap:wrap}
.st-key-panel :is(label[data-baseweb="radio"],[data-testid="stRadioOption"]){background:#fff;border:0.107rem solid var(--line);border-radius:1rem;padding:0.571rem 1.143rem;margin:0}
.st-key-panel :is(label[data-baseweb="radio"]:has(input:checked),[data-testid="stRadioOption"][data-selected="true"]){border-color:var(--pri);background:var(--soft)}

/* ---------- Thẻ tải ảnh (trang chủ) ---------- */
.st-key-upcard{background:linear-gradient(135deg,#e7f8f6 0%,#eef1ff 60%,#f1ecff 100%);border:1px solid #dfe5ff;border-radius:1.571rem;padding:0.857rem 1.286rem;box-shadow:0 0.857rem 2.286rem rgba(79,92,240,.1)}
.uphead{display:flex;align-items:center;gap:1.143rem;margin-bottom:0.571rem}
.uphead .hex{flex:0 0 3.714rem;height:3.714rem;display:grid;place-items:center;background:linear-gradient(135deg,#22c1a5,#3b82f6);clip-path:polygon(50% 0,93% 25%,93% 75%,50% 100%,7% 75%,7% 25%);filter:drop-shadow(0 0.571rem 1rem rgba(34,193,165,.4))}
.uphead h3{margin:0 0 0.286rem;padding:0;font-size:1.25rem;font-weight:800;color:var(--ink)}
.uphead p{margin:0;color:#59618f;font-size:.98rem;line-height:1.5}
.uphead .hw{margin-left:auto;font-family:'Dancing Script',cursive;font-size:1.25rem;color:#3a4690;transform:rotate(-5deg);text-align:right;line-height:1.1;flex:0 0 auto}
.demo{position:relative;display:flex;align-items:center;gap:0.714rem;justify-content:flex-end}
.demo .paper{width:clamp(5rem,8vw,8.571rem);flex:0 1 auto;min-width:0}
.demo .arrow{color:#6d4df0;flex:0 0 auto}
.demo .res{position:relative;flex:1;min-width:0;background:#fff;border-radius:1rem;padding:0.714rem 1rem;box-shadow:0 0.714rem 1.857rem rgba(79,92,240,.14);border:1px solid #e6e9ff;min-width:0}
.demo .res .badge{position:absolute;top:-13px;right:0.714rem;background:#12b886;color:#fff;font-size:.72rem;font-weight:700;border-radius:71.357rem;padding:0.286rem 0.857rem;box-shadow:0 0.429rem 1rem rgba(18,184,134,.35)}
.demo .res h4{margin:0;padding:0;font-size:1.05rem;color:var(--ink)}.demo .res h4 i{float:right;font-style:normal;color:#12b886}
.demo .res b{display:block;color:var(--ink);font-size:.9rem;margin:2px 0 0.571rem}
.demo .res div{font-size:.78rem;color:#4b5380;margin-top:0.286rem;line-height:1.35}
.demo .res small{display:block;margin-top:0.571rem;font-size:.68rem;color:#98a0c7}
[data-testid="stFileUploaderDropzone"]{display:flex;flex-direction:column;align-items:center;justify-content:center;gap:0.571rem;border:2px dashed #8f9dff;background:rgba(255,255,255,.72);border-radius:1.286rem;padding:0.714rem 0.714rem;min-height:9.143rem}
[data-testid="stFileUploaderDropzone"]::before{content:"";order:0;width:3.286rem;height:3.286rem;border-radius:1rem;background:#e7ebff url("__IMGIC__") center/1.714rem no-repeat}
[data-testid="stFileUploaderDropzone"]>span{order:1}
[data-testid="stFileUploaderDropzone"] button *{font-size:0!important}
[data-testid="stFileUploaderDropzone"] button{font-size:0;background:var(--grad);color:#fff;border:0;border-radius:71.357rem;padding:0.571rem 2.143rem;box-shadow:0 0.714rem 1.571rem rgba(79,92,240,.35)}
[data-testid="stFileUploaderDropzone"] button *{color:#fff}
[data-testid="stFileUploaderDropzone"] button p{display:none}
[data-testid="stFileUploaderDropzone"] button::after{content:"Tải ảnh lên";font-size:1rem;font-weight:700;color:#fff}
[data-testid="stFileUploaderDropzone"] button [data-testid="stIconMaterial"]{display:none}
[data-testid="stFileUploaderDropzoneInstructions"]{order:2;text-align:center}
[data-testid="stFileUploaderDropzoneInstructions"] *{font-size:0}
[data-testid="stFileUploaderDropzoneInstructions"] svg,[data-testid="stFileUploaderDropzoneInstructions"] [data-testid="stIconMaterial"]{display:none}
[data-testid="stFileUploaderDropzoneInstructions"]::before{content:"Kéo thả ảnh hoặc chọn file  |  JPG, PNG, WEBP (tối đa 200MB)";font-size:.78rem;color:#6b7299;line-height:1.4}

/* ---------- Khám phá nhanh ---------- */
.sec-title{display:flex;align-items:center;gap:0.714rem;font-weight:700;font-size:1.02rem;color:var(--ink);margin:0.429rem 0.286rem 0}
.sec-title svg{color:var(--pri)}
.tile{display:flex;align-items:center;justify-content:space-between;gap:0.429rem;border-radius:1.286rem;padding:1rem 0.857rem;height:6rem;box-sizing:border-box;border:1px solid rgba(255,255,255,.7);transition:all .18s ease}
.tile .tt{display:flex;flex-direction:column;gap:0.429rem}
.tile .ib{width:2rem;height:2rem;border-radius:50%;background:rgba(255,255,255,.85);display:grid;place-items:center}
.tile b{font-size:.88rem;font-weight:700;line-height:1.25;white-space:pre-line}
.tile .ar{align-self:flex-end;opacity:.85}
.t-pink{background:linear-gradient(135deg,#ffe4f5,#fbd3ee);color:#d6389b}
.t-blue{background:linear-gradient(135deg,#e3eeff,#d4e4ff);color:#2563eb}
.t-green{background:linear-gradient(135deg,#dcf7ea,#c8f1dd);color:#059669}
.t-orange{background:linear-gradient(135deg,#fff1dc,#ffe4bd);color:#d97706}
.t-violet{background:linear-gradient(135deg,#ece8ff,#ddd6ff);color:#6d4df0}

/* ---------- Hành trình + dải tính năng ---------- */
.st-key-journey{background:#fff;border:1px solid var(--line);border-radius:1.286rem;padding:0.571rem 1.286rem;box-shadow:0 0.571rem 1.714rem rgba(79,92,240,.07);margin-top:0}
.jr{display:flex;align-items:center;gap:1.143rem}.jr .tr{font-size:1.9rem;line-height:1}
.jr .tx{flex:1}.jr b{display:block;color:var(--ink);font-size:1.02rem}.jr small{color:var(--mut)}
.jr .bar{flex:1.2;height:0.643rem;border-radius:7.071rem;background:#e8ebff;overflow:hidden;min-width:8.571rem}.jr .bar i{display:block;height:100%;border-radius:7.071rem;background:var(--grad)}
.jr .cnt{font-weight:700;font-size:.88rem;color:var(--ink);white-space:nowrap}
.feat{display:flex;justify-content:space-around;align-items:center;margin:2px 0.286rem 0;color:#4b5380;font-weight:600;font-size:.9rem;flex-wrap:wrap;gap:0.571rem}
.feat span{display:flex;align-items:center;gap:0.571rem}.feat i{width:1px;height:1.286rem;background:#d9defa;display:inline-block}

/* ---------- Cột phải ---------- */
.st-key-quick{background:linear-gradient(135deg,#3b6cf6 0%,#7a52f4 100%);border-radius:1.429rem;padding:0.857rem 1rem 0.714rem;box-shadow:0 1rem 2.143rem rgba(79,92,240,.35)}
.qhead{display:flex;align-items:center;gap:0.857rem;color:#fff;margin-bottom:0.714rem}
.qhead .qi{width:2.714rem;height:2.714rem;border-radius:1rem;background:rgba(255,255,255,.2);display:grid;place-items:center;flex:0 0 auto}
.qhead b{display:block;font-size:1.1rem;color:#fff}.qhead span{font-size:.8rem;color:#e1e6ff}
.st-key-quick [data-testid="stForm"]{border:0;padding:0;background:transparent}
.st-key-quick [data-testid="InputInstructions"]{display:none}
.st-key-quick div[data-baseweb="input"]{background:#fff;border:0;border-radius:71.357rem;height:2.9rem}
.st-key-quick input{color:var(--ink);padding-left:1.1rem}
.st-key-quick [data-testid="stFormSubmitButton"] button{width:2.7rem;height:2.7rem;min-height:0;padding:0;border-radius:50%;background:var(--grad);border:2px solid #fff;color:#fff;font-size:1.15rem;line-height:1}
.st-key-quick [data-testid="stFormSubmitButton"] button p{color:#fff;font-size:1.2rem}
.rcard{background:#fff;border:1px solid var(--line);border-radius:1.286rem;padding:0.714rem 1rem;box-shadow:0 0.571rem 1.714rem rgba(79,92,240,.07);margin-top:0}
.rcard .rh{display:flex;align-items:center;gap:0.714rem;margin-bottom:0.286rem}.rcard .rh svg{color:var(--pri)}
.rcard .rh b{color:var(--ink);font-size:1rem}.rcard .rs{color:var(--mut);font-size:.78rem;margin:0 0 0.429rem 2.429rem}
.st-key-res{background:#fff;border:1px solid var(--line);border-radius:1.286rem;padding:0.714rem 0.857rem 0.286rem;box-shadow:0 0.571rem 1.714rem rgba(79,92,240,.07);margin-top:0}
.st-key-res{gap:0!important}
.st-key-res [data-testid="stElementContainer"],.st-key-res [data-testid="stButton"]{width:100%}
.st-key-res button>div{width:100%;justify-content:flex-start}
.st-key-res button::after{content:"›";margin-left:auto;font-size:1.2rem;color:#9aa3d6}
.st-key-res button{width:100%;justify-content:flex-start;gap:0.714rem;background:transparent;border:0;border-top:1px solid #eef0fb;border-radius:0;padding:.15rem .4rem;min-height:2.1rem;color:var(--ink);box-shadow:none}
.st-key-res button:hover{background:var(--soft);transform:none;box-shadow:none}
.st-key-res button p{font-weight:500;font-size:.9rem;color:var(--ink)}
.st-key-res button [data-testid="stIconMaterial"]{color:var(--pri)}
.hrow{display:flex;align-items:center;gap:0.857rem;padding:0.357rem 2px;border-top:1px solid #eef0fb}
.hrow:first-of-type{border-top:0}
.hrow .hc{width:2.286rem;height:2.286rem;border-radius:50%;display:grid;place-items:center;flex:0 0 auto}
.hrow b{display:block;color:var(--ink);font-size:.88rem}.hrow small{color:var(--mut);font-size:.72rem}
.hrow .hc.c0{background:#d9f7ee;color:#0f9f78}.hrow .hc.c1{background:#ffe0ef;color:#d6389b}.hrow .hc.c2{background:#ffeccd;color:#c2740a}.hrow .hc.c3{background:#e6e0ff;color:#6d4df0}
.rsh{display:flex;align-items:center;gap:0.857rem;color:var(--pri);padding:2px 0.286rem 0.857rem}
.rsh b{display:block;color:var(--ink);font-size:1rem;line-height:1.3}.rsh span{display:block;color:var(--mut);font-size:.78rem;line-height:1.3;margin-top:2px}
.hempty{color:var(--mut);font-size:.85rem;padding:0.571rem 2px}
.st-key-hclear button{border:0;background:transparent;color:var(--pri);font-size:.78rem;padding:0;min-height:0;box-shadow:none}
.quote{display:flex;align-items:center;gap:1rem;border-radius:1.286rem;padding:0.714rem 1rem;margin-top:0;background:linear-gradient(135deg,#f1ecff,#e6f6ff);border:1px solid #e3e6ff}
.quote .lf{font-size:2.2rem;line-height:1}.quote p{margin:0;font-style:italic;color:#3a4690;font-size:.85rem;line-height:1.5}.quote em{display:block;text-align:right;color:var(--pri);font-size:.8rem;font-weight:600;margin-top:0.286rem}

/* ---------- Hộp thoại tài liệu ---------- */
.ptwrap{overflow-x:auto;padding-bottom:0.429rem}
.ptgrid{display:grid;grid-template-columns:repeat(18,minmax(3.857rem,1fr));grid-template-rows:repeat(7,auto) 0.857rem auto auto;gap:0.214rem;min-width:71.429rem}
.ptc{border-radius:0.5rem;padding:0.214rem 0.286rem;text-align:center;line-height:1.15;display:flex;flex-direction:column;align-items:center;min-height:4.429rem;border:1px solid rgba(27,33,80,.08)}
.ptc i{font-style:normal;font-size:.62rem;align-self:flex-start;color:#4b5380}.ptc b{font-size:1.02rem;color:var(--ink)}
.ptc span{font-size:.54rem;color:#3a4172;white-space:nowrap}.ptc em{font-style:normal;font-size:.56rem;color:#4b5380;margin-top:auto}
.ptx{background:#eef0fb;justify-content:center}.ptx b{font-size:.7rem;color:#6b7299}
.ptleg{display:flex;flex-wrap:wrap;gap:0.429rem 1.143rem;margin:0.714rem 0 0.286rem;font-size:.78rem;color:#4b5380}.ptleg span{display:flex;align-items:center;gap:0.429rem}.ptleg i{width:1rem;height:1rem;border-radius:0.286rem;display:inline-block;border:1px solid rgba(27,33,80,.12)}
.kl-kiem{background:#ffd6d6}.kl-kt{background:#ffe3c2}.chuyen{background:#fff0a8}.sau-chuyen{background:#d8f0c8}.a-kim-loai{background:#c7eee3}.phi-kim{background:#c9e6ff}.halogen{background:#d9d4ff}.khi-hiem{background:#f1d2f5}.lantan{background:#ffd0e6}.actini{background:#f4c7d9}
.reftbl{overflow-x:auto}.reftbl table{width:100%;border-collapse:collapse;font-size:.88rem}
.reftbl th{background:#eef1ff;color:var(--ink);text-align:left;padding:0.571rem 0.714rem;position:sticky;top:0}
.reftbl td{padding:0.571rem 0.714rem;border-top:1px solid #eef0fb;color:#2c3466;vertical-align:top}
.rxr{display:flex;gap:0.857rem;align-items:center;padding:0.643rem 2px;border-top:1px solid #eef0fb}.rxr:first-child{border-top:0}
.rxt{flex:0 0 10.714rem;background:#eef1ff;color:#4f5cf0;border-radius:71.357rem;padding:0.214rem 0.714rem;font-size:.74rem;font-weight:700;text-align:center}
.rxe{display:flex;align-items:center;flex-wrap:wrap;gap:0.429rem;color:var(--ink);font-weight:600;font-size:.95rem}
.rxa{display:inline-flex;flex-direction:column;align-items:center;min-width:3.286rem;line-height:1}.rxa small{font-size:.62rem;color:#6b7299;font-weight:500;white-space:nowrap}.rxa i{font-style:normal;font-size:1.2rem;color:#4f5cf0}

/* ---------- Thành phần chung ---------- */
[data-testid="stWidgetLabel"] p{font-weight:700;color:var(--ink);border-left:0.286rem solid var(--pri);padding-left:0.643rem}
.st-key-gsearch [data-testid="stWidgetLabel"],.st-key-quick [data-testid="stWidgetLabel"]{display:none}
button[kind="primary"],[data-testid="stBaseButton-primary"],[data-testid="stBaseButton-primary"]{background:var(--grad);border:0;color:#fff;border-radius:0.857rem;font-weight:700;padding:.55rem 1.3rem}
[data-testid="stBaseButton-secondary"]{border-radius:0.857rem;border-color:var(--line)}
button{transition:transform .15s ease,box-shadow .15s ease}
button:hover{transform:translateY(-1px);box-shadow:0 0.429rem 1.143rem rgba(79,92,240,.22)}
.st-key-jbtn button{border-radius:71.357rem;background:var(--grad);color:#fff;border:0;padding:.6rem 1.4rem;font-weight:700;white-space:nowrap}
.st-key-jbtn button p{color:#fff}
div[data-baseweb="input"],div[data-baseweb="textarea"]{border-radius:0.857rem}
.errtop{color:#d00000;font-weight:800;font-size:1.08rem;line-height:1.5;background:rgba(220,0,0,.08);border:2px solid rgba(220,0,0,.45);border-radius:0.857rem;padding:0.857rem 1.143rem;margin:0.429rem 0 0.857rem}
button:disabled,[data-testid="stBaseButton-primary"]:disabled{background:#dfe3f7!important;color:#8a90b8!important;box-shadow:none!important;transform:none!important}
button:disabled p{color:#8a90b8!important}
.usage{display:flex;align-items:center;gap:0.714rem;background:#fff;border:1px solid var(--line);border-radius:1rem;padding:0.429rem 1rem;color:var(--mut);margin-top:0.429rem;font-size:.85rem}
.usage b{color:var(--pri)}
.disc{color:var(--mut);font-size:.75rem;margin:0.286rem 0.286rem 0}
@media(max-width:900px){.hero3{padding:1.429rem 1.429rem}.hero3 .hr{display:none}.hero3 h1{font-size:1.8rem}.hero3 h1 em{font-size:2.3rem}.st-key-panel{padding:1rem}.uphead .hw{display:none}.demo{flex-wrap:wrap}}
@media(prefers-reduced-motion:reduce){*{animation:none!important;transition:none!important}}
@media(max-width:640px){
html{font-size:14.5px!important}
.block-container{padding:2.8rem .7rem 1rem!important}
.hero3{padding:1rem 1.1rem;border-radius:1.2rem}
.hero3 h1{font-size:1.5rem}.hero3 h1 em{font-size:1.9rem}.hero3 .sub{font-size:.9rem}
.hero3 .chips span{font-size:.78rem;padding:.35rem .7rem}
.st-key-topbar [data-testid="stHorizontalBlock"]{flex-direction:row!important;flex-wrap:nowrap!important;gap:.5rem!important}
.st-key-topbar [data-testid="stColumn"]{min-width:0!important;width:auto!important}
.st-key-topbar [data-testid="stColumn"]:first-child{flex:1 1 auto!important}
.st-key-topbar [data-testid="stColumn"]:last-child{flex:0 0 auto!important}
.tb-right svg{display:none}
.st-key-tilesrow [data-testid="stHorizontalBlock"]{flex-direction:row!important;flex-wrap:wrap!important;gap:.6rem!important}
.st-key-tilesrow [data-testid="stHorizontalBlock"]>[data-testid="stColumn"]{flex:1 1 calc(50% - .6rem)!important;min-width:calc(50% - .6rem)!important}
.uphead .hw{display:none}.uphead{gap:.8rem}.uphead h3{font-size:1.05rem}.uphead p{font-size:.82rem}
.demo{flex-direction:column;align-items:stretch}.demo .paper,.demo .arrow{display:none}
.jr{flex-wrap:wrap}.jr .bar{flex:1 1 100%;order:3}
.feat{display:grid;grid-template-columns:1fr 1fr;gap:.4rem}.feat i{display:none}
.tile{height:5.2rem}
.st-key-panel{padding:.8rem .9rem}
.st-key-quick [data-testid="stHorizontalBlock"]{flex-direction:row!important;flex-wrap:nowrap!important;gap:.5rem!important}
.st-key-quick [data-testid="stColumn"]{min-width:0!important}
.st-key-quick [data-testid="stColumn"]:first-child{flex:1 1 auto!important}
.st-key-quick [data-testid="stColumn"]:last-child{flex:0 0 auto!important;width:auto!important}
.st-key-jbtn button{width:100%}
}
</style>
"""

# Biểu tượng cho từng mục điều hướng (thứ tự = thứ tự trong thanh bên)
NAV_ICON_ORDER = ["home", "structure", "flask", "isomer", "equipment", "mind", "practice", "stats"]
_nav_css = "".join(
    f'.st-key-sidenav label[data-baseweb="radio"]:nth-of-type({i + 1}){{--ic:{_mask(n)}}}.st-key-sidenav [role="radiogroup"]>div:nth-of-type({i + 1}) [data-testid="stRadioOption"]{{--ic:{_mask(n)}}}' for i, n in enumerate(NAV_ICON_ORDER)
)
_ROOT = "clamp(9px, 1.6vh, 16px)" if st.session_state.get("nav", "Trang chủ") == "Trang chủ" else "14px"
CSS = (CSS.replace("__ROOT__", _ROOT).replace("__NAVICONS__", _nav_css)
          .replace("__SEARCHIC__", _uri(icon_svg("search", "#7a82b8", 20)))
          .replace("__IMGIC__", _uri(icon_svg("image", "#4f5cf0", 30))))
st.markdown(CSS, unsafe_allow_html=True)

pw = st.secrets.get("APP_PASSWORD", "")
if pw and st.text_input("Mật khẩu lớp học", type="password") != pw:
    st.stop()

# Khi cập nhật mã, dữ liệu kết quả cũ trong phiên đang mở có thể thiếu trường mới. Xóa các kết quả cũ để tránh lỗi KeyError.
SCHEMA = 5
if st.session_state.get("_schema") != SCHEMA:
    for _k in ("mm", "iso", "iso_nm", "iso_und", "bm", "eq", "rx", "res", "pq"):
        st.session_state.pop(_k, None)
    st.session_state["_schema"] = SCHEMA
st.session_state.setdefault("n", 0)
st.session_state.setdefault("res", None)
st.session_state.setdefault("hist", [])
st.session_state.setdefault("chat", [])


def errbanner(*lines):
    body = "<br>".join(html.escape(str(x)) for x in lines if x)
    return f'<div class="errtop">❌ {body}</div>'


def show_top_error(top, msg, exc=None):
    """Hiện lỗi ở ĐẦU tab, chữ đỏ in đậm."""
    with top:
        st.markdown(errbanner(msg), unsafe_allow_html=True)
        if exc is not None:
            st.caption(f"Chi tiết kỹ thuật: {type(exc).__name__}: {str(exc)[:400]}")


def add_history(label, kind="Công thức"):
    """Ghi một lượt tra cứu vào lịch sử (tối đa 10 mục, bỏ trùng liên tiếp)."""
    h = st.session_state.hist
    label = (label or "").strip()
    if label and not (h and h[0]["q"] == label and h[0]["k"] == kind):
        h.insert(0, {"q": label, "k": kind, "t": datetime.now().strftime("%H:%M")})
        del h[10:]


# ---------------- Công thức cấu tạo viết tay ----------------
def _log_result():
    r = st.session_state.res
    smi = ((r[0].get("smiles") if r else "") or "").strip()
    if smi:
        mol, num = chem_numbers(smi)
        if num:
            add_history(num["formula"], "Công thức")


def _run_structure(top, img_bytes=None, smi=None):
    if st.session_state.n >= MAX_PER_SESSION:
        show_top_error(top, "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục.")
        return
    with st.spinner("Đang phân tích…"):
        try:
            if img_bytes:
                st.session_state.res = (analyze_image_strict(img_bytes), True)
                st.session_state.n += 1
            elif parse_student(smi)[0] is None:
                st.session_state.res = ({"smiles": smi}, False)  # SMILES sai: chấm lỗi ngay, không gọi AI
            else:
                st.session_state.res = (analyze_smiles(smi), False)
                st.session_state.n += 1
            _log_result()
        except Exception as e:
            st.session_state.res = None
            show_top_error(top, "Có lỗi khi phân tích, hãy thử lại hoặc chụp rõ hơn.", e)


def structure_ui():
    top = st.container()  # lỗi và lý do sai hiện ở đầu tab
    # Ảnh hoặc SMILES gửi từ trang chủ / ô Nhập nhanh: phân tích ngay
    pend_img = st.session_state.pop("pending_img", None)
    pend_smi = st.session_state.pop("pending_smi", None)
    if pend_img:
        _run_structure(top, img_bytes=pend_img)
    elif pend_smi:
        _run_structure(top, smi=pend_smi)

    mode = st.radio("Cách nhập", ["🖼️ Tải ảnh lên", "📷 Chụp ảnh", "⌨️ Nhập SMILES"], horizontal=True, key="st_mode")
    img_bytes, smi_in = None, None
    if mode.startswith("📷"):
        f = st.camera_input("Chụp công thức (viết rõ nét, nền sáng, mỗi ảnh một chất)", key="st_cam")
        img_bytes = f.getvalue() if f else None
    elif mode.startswith("🖼️"):
        f = st.file_uploader("Chọn ảnh công thức (viết rõ nét, nền sáng, mỗi ảnh một chất)", type=["jpg", "jpeg", "png", "webp"], key="st_up")
        img_bytes = f.getvalue() if f else None
    else:
        smi_in = st.text_input("Mã SMILES", placeholder="Ví dụ: CC(=O)O", key="st_smi").strip()

    if st.button("Phân tích", type="primary", disabled=not (img_bytes or smi_in), key="st_go"):
        _run_structure(top, img_bytes=img_bytes, smi=smi_in)

    if st.session_state.res:
        r, from_img = st.session_state.res
        smiles_r = (r.get("smiles") or "").strip()
        _, err_r = parse_student(smiles_r) if smiles_r else (None, None)
        if err_r:
            with top:
                st.markdown(errbanner("CÔNG THỨC VẼ SAI: " + err_r), unsafe_allow_html=True)
                im = draw_student(smiles_r)
                if im is not None:
                    st.markdown("**Hình máy đọc lại đúng như hình vẽ (nguyên tử sai khoanh đỏ):**")
                    st.image(im)
                if r.get("luu_y"):
                    st.caption("Ghi chú khi đọc ảnh: " + " • ".join([r["luu_y"]] if isinstance(r["luu_y"], str) else r["luu_y"]))
            st.info("Vì công thức sai hóa trị nên chưa thể gọi tên hay tính khối lượng mol. Hãy sửa lại hình vẽ hoặc sửa SMILES ở ô bên dưới.")
        else:
            st.divider()
            render(r, from_img)
        st.divider()
        fix = st.text_input("Nhận dạng sai? Sửa SMILES rồi phân tích lại", value=r.get("smiles", ""), key="st_fix").strip()
        if st.button("Phân tích lại từ SMILES đã sửa", key="st_fix_go") and fix:
            if st.session_state.n >= MAX_PER_SESSION:
                show_top_error(top, "Đã hết số lượt của phiên này.")
            else:
                with st.spinner("Đang phân tích…"):
                    try:
                        st.session_state.res = (analyze_smiles(fix), False)
                        st.session_state.n += 1
                        _log_result()
                        st.rerun()
                    except Exception as e:
                        show_top_error(top, "Có lỗi khi phân tích lại.")


# ---------------- Trợ lý AI ----------------
ASSIST_RULES = ("Bạn là trợ lý học Hóa học thân thiện cho học sinh Việt Nam (THCS và THPT), bám sát SGK Việt Nam. "
                "Trả lời bằng tiếng Việt, ngắn gọn, rõ ràng; có phương trình, điều kiện, ví dụ khi cần. Dùng danh pháp theo SGK. "
                "Nếu không chắc chắn thì nói rõ và khuyên đối chiếu SGK. Nếu câu hỏi không liên quan đến hóa học, nhẹ nhàng đưa cuộc trò chuyện về hóa học.")


def chat_ai(history):
    contents = [types.Content(role="user" if m["role"] == "user" else "model", parts=[types.Part.from_text(text=m["text"])]) for m in history]
    models = [MODEL]
    fb = st.secrets.get("GEMINI_FALLBACK_MODEL", "gemini-3.5-flash-lite")
    if fb and fb != MODEL:
        models.append(fb)
    last = None
    for m in models:
        for i in range(2):
            try:
                resp = get_client().models.generate_content(
                    model=m, contents=contents, config=types.GenerateContentConfig(system_instruction=ASSIST_RULES))
                return (resp.text or "").strip() or "Mình chưa có câu trả lời, bạn hỏi lại giúp mình nhé."
            except errors.ServerError as e:
                last = e
                time.sleep(2 * (i + 1))
            except errors.ClientError as e:
                last = e
                if getattr(e, "code", None) == 429:
                    time.sleep(3 * (i + 1))
                else:
                    break
    raise last


def assistant_ui():
    chat = st.session_state.chat
    box = st.container()  # tin nhắn nằm trên, ô nhập nằm dưới
    q = st.session_state.pop("pending_chat", None) or st.chat_input("Hỏi bất cứ điều gì về hóa học…", key="chat_in")
    if q:
        chat.append({"role": "user", "text": q})
    with box:
        if not chat:
            st.markdown("Xin chào! Mình là trợ lý hóa học. Bạn cứ hỏi về công thức, phản ứng, danh pháp, đồng phân, cách làm bài tập... nhé.")
        for m in chat:
            with st.chat_message("user" if m["role"] == "user" else "assistant"):
                st.markdown(m["text"])
        if q:
            with st.chat_message("assistant"):
                if st.session_state.n >= MAX_PER_SESSION:
                    ans = "Bạn đã dùng hết số lượt của phiên này. Hãy tải lại trang để tiếp tục."
                    st.markdown(errbanner(ans), unsafe_allow_html=True)
                else:
                    with st.spinner("Đang suy nghĩ…"):
                        try:
                            ans = chat_ai(chat)
                            st.session_state.n += 1
                        except Exception as e:
                            ans = "Có lỗi khi hỏi trợ lý, bạn thử lại sau nhé."
                            st.caption(f"Chi tiết kỹ thuật: {type(e).__name__}: {str(e)[:300]}")
                    st.markdown(ans)
            chat.append({"role": "assistant", "text": ans})
        if chat and st.button("Xóa cuộc trò chuyện", key="chat_clear"):
            st.session_state.chat = []
            st.rerun()


# ---------------- Điều hướng ----------------
HOME = "Trang chủ"
ASSIST = "Trợ lý AI"
SECTIONS = [
    ("Công thức cấu tạo", "Công thức cấu tạo viết tay", "Chụp ảnh công thức vẽ tay kiểu SGK → mã SMILES, tên gọi, khối lượng mol, tính chất.", structure_ui, "structure"),
    ("Phương trình phản ứng", "Phương trình phản ứng", "Đọc phương trình, cho biết loại phản ứng và kiểm tra đã cân bằng chưa.", reaction_ui, "flask"),
    ("Đồng phân", "Đồng phân", "Hỏi bằng tiếng Việt hoặc gõ công thức, máy sinh và đếm đồng phân chính xác.", isomer_ui, "isomer"),
    ("Dụng cụ hóa học", "Dụng cụ hóa học", "Chụp ảnh dụng cụ hoặc bộ thí nghiệm để biết tên, công dụng, cách dùng, an toàn.", equipment_ui, "equipment"),
    ("Sơ đồ tư duy", "Sơ đồ tư duy", "Nhập chủ đề hoặc tên chất, máy vẽ sơ đồ để ôn tập, sửa được và tải về PowerPoint, draw.io.", mindmap_ui, "mind"),
    ("Luyện tập & Chấm bài", "Luyện tập và chấm bài tự động", "Ra đề, nộp bài bằng ảnh vẽ tay, máy chấm và chỉ rõ lỗi sai.", practice_ui, "practice"),
    ("Đánh giá", "Đánh giá độ chính xác", "Dành cho giáo viên và nhóm nghiên cứu: chạy bộ ảnh có đáp án để lấy số liệu.", bench_ui, "stats"),
]
LABELS = [x[0] for x in SECTIONS]
NAV_OPTS = [HOME] + LABELS
STRUCT, REACT, ISO, EQUIP, MIND, PRACT = LABELS[0], LABELS[1], LABELS[2], LABELS[3], LABELS[4], LABELS[5]

st.session_state.setdefault("nav", HOME)
st.session_state.setdefault("nav_side", HOME)


def go(to):
    st.session_state.nav = to
    st.session_state.nav_side = to if to in NAV_OPTS else None


def on_side():
    if st.session_state.nav_side:
        go(st.session_state.nav_side)


def _fold(s):
    s = unicodedata.normalize("NFD", (s or "").lower().replace("đ", "d"))
    return "".join(c for c in s if unicodedata.category(c) != "Mn").strip()


def quick_lookup(text):
    """Đưa chuỗi người dùng nhập tới công cụ phù hợp. Trả về False nếu chuỗi rỗng."""
    t = (text or "").strip()
    if not t:
        return False
    if smiles_like(t):  # giống SMILES: phân tích cấu tạo ngay
        st.session_state.st_mode = "⌨️ Nhập SMILES"
        st.session_state.st_smi = t
        st.session_state.pending_smi = t
        go(STRUCT)
        add_history(t, "Công thức")
    elif is_simple_formula(t) and "C" in t:  # công thức phân tử hữu cơ: đếm đồng phân
        st.session_state.iso_q = t
        go(ISO)
        add_history(t, "Đồng phân")
    else:  # còn lại: hỏi trợ lý AI
        st.session_state.pending_chat = t
        go(ASSIST)
        add_history(t, "Hỏi đáp")
    return True


def on_quick():
    quick_lookup(st.session_state.get("quick_in", ""))
    st.session_state.quick_in = ""


SEARCH_KEYS = {
    STRUCT: "cong thuc cau tao smiles viet tay ctct ten goi khoi luong mol",
    REACT: "phuong trinh phan ung can bang loai phan ung",
    ISO: "dong phan dem dong phan cis trans quang hoc",
    EQUIP: "dung cu hoa hoc thi nghiem ong nghiem binh cau",
    MIND: "so do tu duy mindmap on tap",
    PRACT: "luyen tap cham bai bai tap de kiem tra",
    LABELS[6]: "danh gia do chinh xac benchmark",
    ASSIST: "tro ly ai hoi dap chat",
}


def on_search():
    q = (st.session_state.get("gsearch") or "").strip()
    st.session_state.gsearch = ""
    if not q:
        return
    f = _fold(q)
    for lab, keys in SEARCH_KEYS.items():
        if f in _fold(lab) or (len(f) >= 3 and f in keys):
            go(lab)
            return
    quick_lookup(q)


def on_home_upload():
    f = st.session_state.get("home_up")
    if f is not None:
        st.session_state.pending_img = f.getvalue()
        go(STRUCT)


# ---------------- Tài liệu học tập (dữ liệu tra cứu có sẵn, không gọi AI) ----------------
_PT_NAMES = ("Hiđro,Heli,Liti,Beri,Bo,Cacbon,Nitơ,Oxi,Flo,Neon,Natri,Magie,Nhôm,Silic,Photpho,Lưu huỳnh,Clo,Agon,Kali,Canxi,"
             "Scandi,Titan,Vanađi,Crom,Mangan,Sắt,Coban,Niken,Đồng,Kẽm,Gali,Gecmani,Asen,Selen,Brom,Kripton,Rubiđi,Stronti,Ytri,Zirconi,"
             "Niobi,Molipđen,Tecneti,Rutheni,Rodi,Palađi,Bạc,Cađimi,Inđi,Thiếc,Antimon,Telu,Iot,Xenon,Xesi,Bari,Lantan,Xeri,"
             "Prazeođim,Neođim,Prometi,Samari,Europi,Gađolini,Tecbi,Đisprozi,Honmi,Eribi,Tuli,Ytecbi,Luteti,Hafni,Tantan,Vonfram,"
             "Reni,Osmi,Iriđi,Platin,Vàng,Thủy ngân,Tali,Chì,Bitmut,Poloni,Atatin,Radon,Franxi,Radi,Actini,Thori,Protactini,Urani,"
             "Neptuni,Plutoni,Amerixi,Curi,Berkeli,Califoni,Ensteni,Fecmi,Menđelevi,Nobeli,Lorenxi,Rutherfordi,Đubni,Seaborgi,Bohri,"
             "Hassi,Meitneri,Đacmstati,Rơnghen,Copernixi,Nihoni,Flerovi,Moscovi,Livermori,Tennessin,Oganesson").split(",")
assert len(_PT_NAMES) == 118

_PT_CATS = {
    "kl-kiem": ("Kim loại kiềm", {3, 11, 19, 37, 55, 87}),
    "kl-kt": ("Kim loại kiềm thổ", {4, 12, 20, 38, 56, 88}),
    "chuyen": ("Kim loại chuyển tiếp", set(range(21, 31)) | set(range(39, 49)) | set(range(72, 81)) | set(range(104, 113))),
    "sau-chuyen": ("Kim loại khác", {13, 31, 49, 50, 81, 82, 83, 84, 113, 114, 115, 116}),
    "a-kim-loai": ("Bán kim loại", {5, 14, 32, 33, 51, 52}),
    "phi-kim": ("Phi kim", {1, 6, 7, 8, 15, 16, 34}),
    "halogen": ("Halogen", {9, 17, 35, 53, 85, 117}),
    "khi-hiem": ("Khí hiếm", {2, 10, 18, 36, 54, 86, 118}),
    "lantan": ("Họ lantan", set(range(57, 72))),
    "actini": ("Họ actini", set(range(89, 104))),
}
assert set().union(*[v[1] for v in _PT_CATS.values()]) == set(range(1, 119))
_NO_STABLE = {43, 61} | set(range(84, 90)) | set(range(93, 119))  # không có đồng vị bền: ghi số khối trong ngoặc vuông


def _pt_cat(z):
    return next(k for k, (_, s) in _PT_CATS.items() if z in s)


def _pt_pos(z):
    """Trả về (hàng, cột) trên lưới bảng tuần hoàn dạng 18 cột."""
    if z == 1: return 1, 1
    if z == 2: return 1, 18
    if 3 <= z <= 4: return 2, z - 2
    if 5 <= z <= 10: return 2, z + 8
    if 11 <= z <= 12: return 3, z - 10
    if 13 <= z <= 18: return 3, z
    if 19 <= z <= 36: return 4, z - 18
    if 37 <= z <= 54: return 5, z - 36
    if z in (55, 56): return 6, z - 54
    if 57 <= z <= 71: return 9, z - 54
    if 72 <= z <= 86: return 6, z - 68
    if z in (87, 88): return 7, z - 86
    if 89 <= z <= 103: return 10, z - 86
    return 7, z - 100


def _pt_mass(z):
    m = Chem.GetPeriodicTable().GetAtomicWeight(z)
    return f"[{round(m)}]" if z in _NO_STABLE else f"{m:.2f}"


@st.dialog("Bảng tuần hoàn các nguyên tố hóa học", width="large")
def periodic_dialog():
    pt = Chem.GetPeriodicTable()
    cells = ['<div class="ptc ptx" style="grid-row:6;grid-column:3"><b>57–71</b></div>',
             '<div class="ptc ptx" style="grid-row:7;grid-column:3"><b>89–103</b></div>']
    for z in range(1, 119):
        r, c = _pt_pos(z)
        sym, nm = pt.GetElementSymbol(z), _PT_NAMES[z - 1]
        cells.append(f'<div class="ptc {_pt_cat(z)}" style="grid-row:{r};grid-column:{c}" title="{z} · {nm}">'
                     f'<i>{z}</i><b>{sym}</b><span>{nm}</span><em>{_pt_mass(z)}</em></div>')
    legend = "".join(f'<span><i class="{k}"></i>{v[0]}</span>' for k, v in _PT_CATS.items())
    st.markdown(f'<div class="ptwrap"><div class="ptgrid">{"".join(cells)}</div></div><div class="ptleg">{legend}</div>', unsafe_allow_html=True)
    st.caption("Khối lượng nguyên tử tính theo RDKit (đơn vị amu). Số trong ngoặc vuông [ ] là số khối của đồng vị bền nhất, "
               "dùng cho nguyên tố không có đồng vị bền. Hãy đối chiếu bảng trong SGK khi làm bài.")
    opts = [f"{z} · {pt.GetElementSymbol(z)} · {_PT_NAMES[z - 1]}" for z in range(1, 119)]
    pick = st.selectbox("Tra cứu một nguyên tố", opts, index=5, key="pt_pick")
    z = opts.index(pick) + 1
    r, c = _pt_pos(z)
    a, b, d, e = st.columns(4)
    a.metric("Số hiệu (Z)", z)
    b.metric("Ký hiệu", pt.GetElementSymbol(z))
    d.metric("Khối lượng", _pt_mass(z))
    e.metric("Nhóm", _PT_CATS[_pt_cat(z)][0])
    if r >= 9:
        st.markdown(f"**Chu kì:** {6 if r == 9 else 7} (hàng riêng bên dưới bảng)")
    else:
        st.markdown(f"**Chu kì:** {r} · **Cột:** {c}")


def _sub(s):
    """Chuyển chữ số sau chữ cái hoặc dấu ngoặc đóng thành chỉ số dưới."""
    return re.sub(r"(?<=[A-Za-z\)])(\d+)", lambda m: m.group(1).translate(SUB), s).replace(")n", ")ₙ")


ORG_ROWS = [  # (loại hợp chất, công thức chung, nhóm chức hoặc đặc điểm, ví dụ)
    ("Ankan", "C<sub>n</sub>H<sub>2n+2</sub> (n ≥ 1)", "Chỉ có liên kết đơn", "CH4 (metan), C2H6 (etan)"),
    ("Xicloankan (1 vòng)", "C<sub>n</sub>H<sub>2n</sub> (n ≥ 3)", "Mạch vòng, liên kết đơn", "C3H6 (xiclopropan)"),
    ("Anken", "C<sub>n</sub>H<sub>2n</sub> (n ≥ 2)", "1 liên kết đôi C=C", "C2H4 (eten)"),
    ("Ankađien", "C<sub>n</sub>H<sub>2n–2</sub> (n ≥ 3)", "2 liên kết đôi C=C", "C4H6 (buta-1,3-đien)"),
    ("Ankin", "C<sub>n</sub>H<sub>2n–2</sub> (n ≥ 2)", "1 liên kết ba C≡C", "C2H2 (etin)"),
    ("Aren (đồng đẳng benzen)", "C<sub>n</sub>H<sub>2n–6</sub> (n ≥ 6)", "Có vòng benzen", "C6H6 (benzen), C7H8 (toluen)"),
    ("Dẫn xuất halogen no, đơn chức", "C<sub>n</sub>H<sub>2n+1</sub>X", "–X (F, Cl, Br, I)", "CH3Cl, C2H5Br"),
    ("Ancol no, đơn chức, mạch hở", "C<sub>n</sub>H<sub>2n+1</sub>OH (C<sub>n</sub>H<sub>2n+2</sub>O)", "–OH gắn vào C no", "CH3OH, C2H5OH"),
    ("Phenol", "C<sub>6</sub>H<sub>5</sub>OH", "–OH gắn trực tiếp vào vòng benzen", "C6H5OH"),
    ("Anđehit no, đơn chức, mạch hở", "C<sub>n</sub>H<sub>2n+1</sub>CHO (C<sub>n</sub>H<sub>2n</sub>O)", "–CHO", "HCHO, CH3CHO"),
    ("Xeton no, đơn chức, mạch hở", "C<sub>n</sub>H<sub>2n</sub>O (n ≥ 3)", ">C=O", "CH3COCH3 (propanon)"),
    ("Axit cacboxylic no, đơn chức, mạch hở", "C<sub>n</sub>H<sub>2n+1</sub>COOH (C<sub>n</sub>H<sub>2n</sub>O<sub>2</sub>)", "–COOH", "HCOOH, CH3COOH"),
    ("Este no, đơn chức, mạch hở", "C<sub>n</sub>H<sub>2n</sub>O<sub>2</sub> (n ≥ 2)", "–COO–", "HCOOCH3, CH3COOC2H5"),
    ("Amin no, đơn chức, mạch hở", "C<sub>n</sub>H<sub>2n+3</sub>N (n ≥ 1)", "–NH2 (amin bậc 1)", "CH3NH2, C2H5NH2"),
    ("Amino axit", "H<sub>2</sub>N–R–COOH", "–NH2 và –COOH", "H2NCH2COOH (glyxin)"),
    ("Glucozơ, fructozơ", "C<sub>6</sub>H<sub>12</sub>O<sub>6</sub>", "Monosaccarit", "C6H12O6"),
    ("Saccarozơ, mantozơ", "C<sub>12</sub>H<sub>22</sub>O<sub>11</sub>", "Đisaccarit", "C12H22O11"),
    ("Tinh bột, xenlulozơ", "(C<sub>6</sub>H<sub>10</sub>O<sub>5</sub>)<sub>n</sub>", "Polisaccarit", "(C6H10O5)n"),
    ("Chất béo", "(RCOO)<sub>3</sub>C<sub>3</sub>H<sub>5</sub>", "Trieste của glixerol với axit béo", "(C17H35COO)3C3H5 (tristearin)"),
]


@st.dialog("Công thức hợp chất hữu cơ", width="large")
def organic_dialog():
    rows = "".join(f"<tr><td><b>{a}</b></td><td>{b}</td><td>{_sub(c)}</td><td>{_sub(d)}</td></tr>" for a, b, c, d in ORG_ROWS)
    st.markdown('<div class="reftbl"><table><thead><tr><th>Loại hợp chất</th><th>Công thức chung</th><th>Nhóm chức / đặc điểm</th><th>Ví dụ</th></tr></thead>'
                f'<tbody>{rows}</tbody></table></div>', unsafe_allow_html=True)
    st.markdown("**Độ bất bão hòa:** k = (2C + 2 + N − H − X) / 2, trong đó X là tổng số nguyên tử halogen. "
                "k bằng tổng số vòng và số liên kết π (một liên kết đôi tính 1, một liên kết ba tính 2).")
    st.caption("Các chất đồng đẳng hơn kém nhau một hoặc nhiều nhóm CH2. Hãy đối chiếu SGK khi làm bài.")
    st.button("Mở công cụ Công thức cấu tạo", key="org_open", on_click=go, args=(STRUCT,))


# Mỗi phương trình: (loại, vế trái, vế phải, điều kiện, mũi tên). Vế là danh sách (hệ số, công thức).
RX_ORG = [
    ("Thế", [(1, "CH4"), (1, "Cl2")], [(1, "CH3Cl"), (1, "HCl")], "ánh sáng", "→"),
    ("Cộng H2", [(1, "CH2=CH2"), (1, "H2")], [(1, "CH3CH3")], "Ni, t°", "→"),
    ("Cộng Br2", [(1, "CH2=CH2"), (1, "Br2")], [(1, "CH2BrCH2Br")], "", "→"),
    ("Cộng HX", [(1, "CH2=CH2"), (1, "HCl")], [(1, "CH3CH2Cl")], "", "→"),
    ("Cộng H2O", [(1, "CH2=CH2"), (1, "H2O")], [(1, "CH3CH2OH")], "H⁺, t°", "→"),
    ("Cộng Br2 (ankin)", [(1, "CH≡CH"), (2, "Br2")], [(1, "CHBr2CHBr2")], "", "→"),
    ("Tách nước", [(1, "CH3CH2OH")], [(1, "CH2=CH2"), (1, "H2O")], "H2SO4 đặc, 170°C", "→"),
    ("Oxi hóa ancol", [(2, "CH3CH2OH"), (1, "O2")], [(2, "CH3CHO"), (2, "H2O")], "Cu, t°", "→"),
    ("Đốt cháy", [(1, "C2H4"), (3, "O2")], [(2, "CO2"), (2, "H2O")], "t°", "→"),
    ("Đốt cháy", [(2, "C2H6"), (7, "O2")], [(4, "CO2"), (6, "H2O")], "t°", "→"),
    ("Thế (benzen)", [(1, "C6H6"), (1, "Br2")], [(1, "C6H5Br"), (1, "HBr")], "Fe, t°", "→"),
    ("Nitro hóa", [(1, "C6H6"), (1, "HNO3")], [(1, "C6H5NO2"), (1, "H2O")], "H2SO4 đặc", "→"),
    ("Thế (phenol)", [(1, "C6H5OH"), (3, "Br2")], [(1, "C6H2Br3OH↓"), (3, "HBr")], "", "→"),
    ("Este hóa", [(1, "CH3COOH"), (1, "C2H5OH")], [(1, "CH3COOC2H5"), (1, "H2O")], "H2SO4 đặc, t°", "⇌"),
    ("Xà phòng hóa", [(1, "CH3COOC2H5"), (1, "NaOH")], [(1, "CH3COONa"), (1, "C2H5OH")], "t°", "→"),
    ("Tráng bạc", [(1, "CH3CHO"), (2, "AgNO3"), (3, "NH3"), (1, "H2O")], [(1, "CH3COONH4"), (2, "Ag↓"), (2, "NH4NO3")], "t°", "→"),
    ("Lên men", [(1, "C6H12O6")], [(2, "C2H5OH"), (2, "CO2↑")], "enzim, 30–35°C", "→"),
    ("Axit – bazơ", [(1, "CH3COOH"), (1, "NaOH")], [(1, "CH3COONa"), (1, "H2O")], "", "→"),
    ("Amin + axit", [(1, "C6H5NH2"), (1, "HCl")], [(1, "C6H5NH3Cl")], "", "→"),
]
RX_INORG = [
    ("Kim loại + axit", [(1, "Zn"), (2, "HCl")], [(1, "ZnCl2"), (1, "H2↑")], "", "→"),
    ("Kim loại + axit", [(2, "Al"), (6, "HCl")], [(2, "AlCl3"), (3, "H2↑")], "", "→"),
    ("Kim loại + muối", [(1, "Fe"), (1, "CuSO4")], [(1, "FeSO4"), (1, "Cu")], "", "→"),
    ("Kim loại + nước", [(2, "Na"), (2, "H2O")], [(2, "NaOH"), (1, "H2↑")], "", "→"),
    ("Trung hòa", [(1, "NaOH"), (1, "HCl")], [(1, "NaCl"), (1, "H2O")], "", "→"),
    ("Trao đổi (kết tủa)", [(1, "AgNO3"), (1, "NaCl")], [(1, "AgCl↓"), (1, "NaNO3")], "", "→"),
    ("Trao đổi (kết tủa)", [(1, "BaCl2"), (1, "H2SO4")], [(1, "BaSO4↓"), (2, "HCl")], "", "→"),
    ("Muối cacbonat + axit", [(1, "CaCO3"), (2, "HCl")], [(1, "CaCl2"), (1, "CO2↑"), (1, "H2O")], "", "→"),
    ("Oxit axit + bazơ", [(1, "Ca(OH)2"), (1, "CO2")], [(1, "CaCO3↓"), (1, "H2O")], "", "→"),
    ("Phân hủy", [(1, "CaCO3")], [(1, "CaO"), (1, "CO2↑")], "t°", "→"),
    ("Phân hủy", [(2, "KClO3")], [(2, "KCl"), (3, "O2↑")], "MnO2, t°", "→"),
    ("Hóa hợp", [(2, "H2"), (1, "O2")], [(2, "H2O")], "t°", "→"),
    ("Hóa hợp", [(2, "Mg"), (1, "O2")], [(2, "MgO")], "t°", "→"),
    ("Oxit kim loại + H2", [(1, "CuO"), (1, "H2")], [(1, "Cu"), (1, "H2O")], "t°", "→"),
    ("Luyện gang", [(1, "Fe2O3"), (3, "CO")], [(2, "Fe"), (3, "CO2")], "t°", "→"),
    ("Axit đặc", [(1, "Cu"), (2, "H2SO4")], [(1, "CuSO4"), (1, "SO2↑"), (2, "H2O")], "H2SO4 đặc, t°", "→"),
]


def _atoms(formula):
    """Đếm nguyên tử của một công thức (hỗ trợ ngoặc), bỏ ký hiệu ↑ ↓ = ≡."""
    f = re.sub(r"[↑↓=≡]", "", formula)
    stack = [{}]
    for tok in re.findall(r"([A-Z][a-z]?)(\d*)|(\()|\)(\d*)", f):
        el, n, op, cn = tok
        if el:
            stack[-1][el] = stack[-1].get(el, 0) + (int(n) if n else 1)
        elif op:
            stack.append({})
        else:
            grp = stack.pop()
            for k, v in grp.items():
                stack[-1][k] = stack[-1].get(k, 0) + v * (int(cn) if cn else 1)
    return stack[0]


def rx_balanced(left, right):
    def tot(side):
        c = Counter()
        for k, f in side:
            for el, v in _atoms(f).items():
                c[el] += k * v
        return c
    return tot(left) == tot(right)


def _rx_html(rows):
    out = []
    for typ, L, R, cond, arrow in rows:
        side = lambda s: " + ".join((f"{k} " if k > 1 else "") + html.escape(_sub(f)) for k, f in s)
        arr = f'<span class="rxa"><small>{html.escape(_sub(cond))}</small><i>{arrow}</i></span>'
        out.append(f'<div class="rxr"><span class="rxt">{html.escape(_sub(typ))}</span><div class="rxe">{side(L)}{arr}{side(R)}</div></div>')
    return "".join(out)


@st.dialog("Phương trình phản ứng thường gặp", width="large")
def reactions_dialog():
    t1, t2 = st.tabs(["Hóa hữu cơ", "Hóa vô cơ"])
    with t1:
        st.markdown(_rx_html(RX_ORG), unsafe_allow_html=True)
        st.caption("Trùng hợp: nCH₂=CH₂ → (–CH₂–CH₂–)ₙ (điều kiện: t°, p, xúc tác).")
    with t2:
        st.markdown(_rx_html(RX_INORG), unsafe_allow_html=True)
    st.caption("Các phương trình đã được máy kiểm tra cân bằng nguyên tử. ↑ là chất khí, ↓ là chất kết tủa. Điều kiện ghi trên mũi tên.")
    st.button("Mở công cụ Phương trình phản ứng", key="rx_open", on_click=go, args=(REACT,))


@st.dialog("Mẹo & lưu ý", width="large")
def tips_dialog():
    st.markdown("**Chụp ảnh để máy đọc chính xác**\n"
                "- Viết công thức rõ nét, nền sáng, mỗi ảnh chỉ một chất.\n"
                "- Chụp thẳng góc, đủ sáng, tránh bóng đổ và nếp gấp giấy.\n"
                "- Công thức sai hóa trị sẽ được máy chỉ rõ nguyên tử sai, hãy sửa lại hình vẽ rồi chụp lại.\n"
                "- Có thể sửa mã SMILES rồi phân tích lại ngay dưới kết quả.\n\n"
                "**Gọi tên và quy tắc hay dùng (theo SGK)**\n"
                "- Danh pháp thay thế: chọn mạch chính dài nhất có nhóm chức, đánh số từ phía gần nhóm chức, "
                "gọi tên nhánh theo thứ tự chữ cái.\n"
                "- Quy tắc Mác-côp-nhi-côp: khi cộng HX hoặc H₂O vào anken bất đối xứng, H cộng vào C có nhiều H hơn, "
                "X hoặc OH cộng vào C có ít H hơn.\n"
                "- Quy tắc Zai-xép: khi tách HX hoặc H₂O, ưu tiên tách H ở C bên cạnh có bậc cao hơn.\n"
                "- Thế vào vòng benzen: –OH, –NH₂, –CH₃ định hướng vị trí o- và p-; –NO₂, –COOH, –CHO định hướng vị trí m-.\n"
                "- Khi cân bằng, kiểm tra số nguyên tử từng nguyên tố hai vế bằng nhau; C luôn có hóa trị 4.\n\n"
                "**Lưu ý**\n"
                "- Kết quả do AI gợi ý có thể sai ở nét khó đọc: luôn đối chiếu SGK.")


# ---------------- Các khối giao diện ----------------
def topbar():
    with st.container(key="topbar"):
        a, b = st.columns([5, 2], vertical_alignment="center")
        with a:
            st.text_input("Tìm kiếm", key="gsearch", placeholder="Tìm công cụ, công thức, tên chất…", on_change=on_search, label_visibility="collapsed")
        with b:
            st.markdown(f'<div class="tb-right">{icon_svg("sun", "#5b6293", 22)}{icon_svg("bell", "#5b6293", 22)}<div class="av">T</div></div>',
                        unsafe_allow_html=True)


def hero():
    st.markdown(
        '<div class="hero3"><div class="hl"><span class="pill">Chào bạn ! 👋</span>'
        f'<h1>Khám phá thế giới Hóa học<br>cùng <em>{html.escape(APP_NAME)}</em></h1>'
        '<div class="sub"><span>Nhanh chóng</span><i>•</i><span>Chính xác</span><i>•</i><span>Thông minh</span></div>'
        '<div class="chips"><span>⚡ Nhanh chóng</span><span>🎯 Chính xác</span><span>💡 Tiện lợi</span><span>💜 Học mọi lúc</span></div></div>'
        f'<div class="hr"><img class="art" src="{_uri(HERO_ART)}"><div class="slogan">Mỗi phản ứng<br>là một khám phá!<u></u></div></div></div>',
        unsafe_allow_html=True,
    )


def page_banner(icon, title, desc):
    st.markdown(f'<div class="pbanner"><div class="pb-ic">{icon_svg(icon, "#fff", 28)}</div><div><b>{html.escape(title)}</b>'
                f'<span>{html.escape(desc)}</span></div></div>', unsafe_allow_html=True)


TILES = [("Công thức\ncấu tạo", "structure", "t-pink", STRUCT), ("Phương trình\nphản ứng", "flask", "t-blue", REACT),
         ("Đồng phân", "isomer", "t-green", ISO), ("Dụng cụ\nhóa học", "equipment", "t-orange", EQUIP), ("Sơ đồ tư duy", "mind", "t-violet", MIND)]
FEATS = [("⚡", "Nhanh chóng"), ("🎯", "Chính xác"), ("💡", "Tiện lợi"), ("💜", "Học mọi lúc")]
HKIND = {"Công thức": ("structure", "c1"), "Phản ứng": ("flask", "c3"), "Đồng phân": ("isomer", "c0"), "Dụng cụ": ("equipment", "c2"),
         "Sơ đồ": ("mind", "c3"), "Luyện tập": ("practice", "c0"), "Hỏi đáp": ("chat", "c1")}
GOAL = 15


def home_ui():
    left, right = st.columns([2.15, 1], gap="large")
    with left:
        with st.container(key="upcard"):
            st.markdown(
                f'<div class="uphead"><div class="hex">{icon_svg("pen", "#fff", 34, 2)}</div><div><h3>Công thức cấu tạo viết tay</h3>'
                '<p>Chụp ảnh công thức vẽ tay kiểu SGK → mã SMILES, tên gọi, khối lượng mol, tính chất.</p></div>'
                f'<div class="hw">Chụp nhanh<br>Nhận kết quả ngay!</div></div>', unsafe_allow_html=True)
            cu, cd = st.columns([1, 1.5], gap="medium", vertical_alignment="center")
            with cu:
                st.file_uploader("Tải ảnh công thức", type=["jpg", "jpeg", "png", "webp"], key="home_up",
                                 label_visibility="collapsed", on_change=on_home_upload)
            with cd:
                st.markdown(
                    f'<div class="demo"><img class="paper" src="{_uri(PAPER_SVG)}"><span class="arrow">{icon_svg("arrow", "#6d4df0", 28, 2)}</span>'
                    '<div class="res"><span class="badge">✦ Kết quả ngay lập tức</span><h4>C₇H₆O₂<i>●</i></h4><b>Axit benzoic</b>'
                    '<div>SMILES: O=C(O)c1ccccc1</div><div>Khối lượng mol: 122.12 g/mol</div><div>Tính chất: Axit yếu, rắn, ít tan trong nước</div>'
                    '<small>Ví dụ minh họa</small></div></div>', unsafe_allow_html=True)

        st.markdown(f'<div class="sec-title">{icon_svg("home", "currentColor", 22)}Khám phá nhanh</div>', unsafe_allow_html=True)
        with st.container(key="tilesrow"):
            for c, (txt, ic, cls, target), i in zip(st.columns(5, gap="small"), TILES, range(5)):
                with c, st.container(key=f"qk_{i}"):
                    st.markdown(f'<div class="tile {cls}"><div class="tt"><div class="ib">{icon_svg(ic, "currentColor", 18, 2)}</div><b>{txt}</b></div>'
                                f'<span class="ar">{icon_svg("arrow", "currentColor", 18, 2.2)}</span></div>', unsafe_allow_html=True)
                    st.button("Mở", key=f"qkb_{i}", on_click=go, args=(target,))

        done = len(st.session_state.get("ph", []))
        pct = min(100, round(100 * done / GOAL))
        with st.container(key="journey"):
            jc, jb = st.columns([4.2, 1], vertical_alignment="center")
            with jc:
                st.markdown(
                    f'<div class="jr"><div class="tr">🏆</div><div class="tx"><b>Hành trình học tập của bạn</b></div>'
                    f'<div class="bar"><i style="width:{pct}%"></i></div><div class="cnt">{done}/{GOAL} bài hoàn thành</div></div>', unsafe_allow_html=True)
            with jb:
                with st.container(key="jbtn"):
                    st.button("Tiếp tục học  ›", key="jgo", on_click=go, args=(PRACT,))
        st.markdown('<div class="feat">' + '<i></i>'.join(f"<span>{e} {t}</span>" for e, t in FEATS) + '</div>', unsafe_allow_html=True)

    with right:
        with st.container(key="quick"):
            st.markdown(f'<div class="qhead"><div class="qi">{icon_svg("pen", "#fff", 24, 2)}</div><div><b>Nhập nhanh</b>'
                        '<span>Nhập công thức để tra cứu ngay</span></div></div>', unsafe_allow_html=True)
            with st.form("quick_form", border=False, clear_on_submit=False):
                qa, qb = st.columns([5, 1.2], vertical_alignment="center", gap="small")
                with qa:
                    st.text_input("Công thức", key="quick_in", placeholder="VD: C4H10O, CCO, axit axetic…", label_visibility="collapsed")
                with qb:
                    st.form_submit_button("➜", on_click=on_quick)

        with st.container(key="res"):
            st.markdown(f'<div class="rsh">{icon_svg("equipment", "currentColor", 24)}'
                        '<div><b>Tài liệu học tập</b><span>Kho tài liệu &amp; công thức quan trọng</span></div></div>', unsafe_allow_html=True)
            if st.button("Bảng tuần hoàn", icon=":material/grid_view:", key="r_pt"):
                periodic_dialog()
            if st.button("Công thức hữu cơ", icon=":material/hub:", key="r_org"):
                organic_dialog()
            if st.button("Phương trình phản ứng", icon=":material/science:", key="r_rx"):
                reactions_dialog()
            if st.button("Mẹo & lưu ý", icon=":material/lightbulb:", key="r_tip"):
                tips_dialog()

        hist = st.session_state.hist
        rows = "".join(
            f'<div class="hrow"><div class="hc {HKIND.get(h.get("k"), HKIND["Công thức"])[1]}">{icon_svg(HKIND.get(h.get("k"), HKIND["Công thức"])[0], "currentColor", 20, 2)}</div>'
            f'<div><b>{html.escape(h["q"][:40] + ("…" if len(h["q"]) > 40 else ""))}</b><small>{html.escape(h.get("k", ""))} · Hôm nay · {h["t"]}</small></div></div>'
            for h in hist[:6]) or '<div class="hempty">Chưa có lượt tra cứu nào. Hãy thử Nhập nhanh hoặc dùng một công cụ bên trái.</div>'
        st.markdown(f'<div class="rcard"><div class="rh">{icon_svg("search", "currentColor", 22)}<b>Lịch sử tra cứu</b></div>{rows}</div>', unsafe_allow_html=True)
        if hist:
            with st.container(key="hclear"):
                st.button("Xóa lịch sử", key="hist_clear", on_click=lambda: st.session_state.hist.clear())
        st.markdown(f'<div class="quote"><div class="lf">🌿</div><p>“Hóa học không khó, chỉ cần bạn có công cụ đúng!”<em>— {html.escape(APP_NAME)}</em></p></div>',
                    unsafe_allow_html=True)


with st.sidebar:
    st.markdown(f'<div class="sb-brand"><img src="{_uri(LOGO_SVG)}"><div><b>{html.escape(APP_NAME[:3])}<i>{html.escape(APP_NAME[3:])}</i></b>'
                '<span>AI trong tay – Hóa học dễ dàng</span></div></div>', unsafe_allow_html=True)
    with st.container(key="sidenav"):
        st.radio("Điều hướng", NAV_OPTS, index=None, key="nav_side", on_change=on_side, label_visibility="collapsed")
    with st.container(key="aibtn"):
        st.markdown(f'<div class="aicard"><img src="{_uri(AI_FACE)}"><div><b>Trợ lý AI</b><span>Hỏi bất cứ điều gì về<br>hóa học, mình luôn ở đây!</span></div>'
                    '<div class="go">›</div></div>', unsafe_allow_html=True)
        st.button("Mở trợ lý AI", key="aibtn_b", on_click=go, args=(ASSIST,))
    st.markdown(f'<div class="mascot"><img src="{_uri(MASCOT_SVG)}"><div class="hw">Học hóa<br>thật dễ!</div></div>', unsafe_allow_html=True)

# Trên điện thoại: tự đóng thanh bên sau khi chọn một mục (mã chạy trong trang chính để không bị hủy khi làm mới)
components.html("""<script>
(function(){const w=window.parent,d=w.document;if(w.__molsnapNav)return;w.__molsnapNav=true;
const s=d.createElement('script');
s.textContent="document.addEventListener('click',function(e){if(window.innerWidth>640)return;"
+"var o=e.target.closest('[data-testid=\\"stRadioOption\\"],.st-key-aibtn');"
+"if(o&&o.closest('[data-testid=\\"stSidebar\\"]')){setTimeout(function(){"
+"var b=document.querySelector('[data-testid=\\"stSidebarCollapseButton\\"] button');if(b)b.click();},350);}},true);";
d.head.appendChild(s);})();
</script>""", height=0)

topbar()

nav = st.session_state.nav
if nav == HOME or (nav not in LABELS and nav != ASSIST):
    hero()
    home_ui()
elif nav == ASSIST:
    page_banner("chat", "Trợ lý AI", "Hỏi bất cứ điều gì về hóa học, mình luôn ở đây!")
    with st.container(key="panel"):
        assistant_ui()
else:
    label, title, desc, fn, ic = SECTIONS[LABELS.index(nav)]
    page_banner(ic, title, desc)
    with st.container(key="panel"):
        fn()

st.markdown(f'<div class="usage">📄 Lượt phân tích đã dùng trong phiên này: <b>{st.session_state.n}/{MAX_PER_SESSION}</b></div>'
            '<div class="disc">ⓘ Kết quả do AI gợi ý, có thể sai ở nét viết khó đọc. Công thức, khối lượng mol, cân bằng và số đồng phân do bộ máy hóa học tính. Hãy đối chiếu SGK.</div>',
            unsafe_allow_html=True)
